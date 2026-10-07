# -*- coding: utf-8 -*-
import re
from docx import Document
from docx.oxml.ns import qn

SRC = r"c:\Users\HEFEI\.trae-cn\attachments\6ac5f18c3e706c34f9ec0610\3597e0cd-ad7f-4049-8fc5-b5679b0ff62e_装饰装修专项施工方案_20261007_第4轮.docx"
doc = Document(SRC)

# 1. 内嵌图片数量（按文档正文顺序遍历 drawing）
body = doc.element.body
blips = body.findall('.//' + qn('a:blip'))
print("内嵌图片数(a:blip):", len(blips))

# 2. 图题、表题序列
fig_pat = re.compile(r'^图\s*(\d+)[-－.](\d+)')
tab_pat = re.compile(r'^表\s*(\d+)[-－.](\d+)')
figs, tabs = [], []
for p in doc.paragraphs:
    t = p.text.strip()
    m = fig_pat.match(t)
    if m: figs.append((int(m.group(1)), int(m.group(2)), t))
    m2 = tab_pat.match(t)
    if m2: tabs.append((int(m2.group(1)), int(m2.group(2)), t))
print("\n图题序列:")
for f in figs: print("  ", f[2])
print("表题序列(段落中):")
for t in tabs: print("  ", t[2])

# 3. 表格内容速览：前两行 + 行列数
print("\n表格清单:")
for i, tbl in enumerate(doc.tables):
    rows = len(tbl.rows); cols = len(tbl.columns)
    first = " | ".join(c.text.strip()[:14] for c in tbl.rows[0].cells[:min(cols, 6)])
    second = ""
    if rows > 1:
        second = " || ".join(c.text.strip()[:14] for c in tbl.rows[1].cells[:min(cols, 6)])
    print(f"  T{i:02d} {rows}行x{cols}列 | 表头: {first}")
    if second:
        print(f"        首行: {second}")

# 4. 占位符/待补充/未闭合围栏/双空格断字
ph_pat = re.compile(r'(【[^】]{0,20}】|待补充|待定|TODO|XXX|xxx|\{\{[^}]{0,20}\}\}|_{3,}|（\s*）)')
print("\n占位/空缺标记:")
cnt = 0
for idx, p in enumerate(doc.paragraphs):
    for m in ph_pat.finditer(p.text):
        s = max(0, m.start()-18); e = min(len(p.text), m.end()+18)
        print(f"  段{idx}: ...{p.text[s:e]}...")
        cnt += 1
        if cnt > 60: break
    if cnt > 60: break
print("占位标记总数(上限统计):", cnt)

# 连续双空格（公式变量丢失特征）
dbl = [(i, p.text) for i, p in enumerate(doc.paragraphs) if "  " in p.text]
print("\n含连续空格的段落数:", len(dbl))
for i, t in dbl[:25]:
    pos = [m.start() for m in re.finditer(r'  +', t)]
    snip = t[:110]
    print(f"  段{i}: {snip}")

# 5. 英文短语混入（连续3个以上英文单词）
en_pat = re.compile(r'[A-Za-z]+(?:\s+[A-Za-z]+){2,}')
print("\n英文短语混入:")
seen = 0
for i, p in enumerate(doc.paragraphs):
    for m in en_pat.finditer(p.text):
        s = max(0, m.start()-15); e = min(len(p.text), m.end()+15)
        print(f"  段{i}: ...{p.text[s:e]}...")
        seen += 1
        if seen > 25: break
    if seen > 25: break

# 6. 正文中“如图/见表”引用与图号对应
refs = re.findall(r'图\s*(\d+[-－.]\d+)', "\n".join(p.text for p in doc.paragraphs))
print("\n正文图号引用(含图题):", sorted(set(refs)))
