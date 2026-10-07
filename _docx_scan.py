# -*- coding: utf-8 -*-
import re, collections
from docx import Document
from docx.oxml.ns import qn

SRC = r"c:\Users\HEFEI\.trae-cn\attachments\6ac5f18c3e706c34f9ec0610\3597e0cd-ad7f-4049-8fc5-b5679b0ff62e_装饰装修专项施工方案_20261007_第4轮.docx"

doc = Document(SRC)

# ---------- 1. 段落总览（样式 / 编号 / 文本 / 斜体） ----------
paras = doc.paragraphs
print("段落总数:", len(paras), " 表格数:", len(doc.tables))
style_counter = collections.Counter(p.style.name if p.style else "?" for p in paras)
print("样式分布:", dict(style_counter))

# 识别大纲级别
def outline_level(p):
    pPr = p._p.find(qn('w:pPr'))
    if pPr is None:
        return None
    ol = pPr.find(qn('w:outlineLvl'))
    if ol is not None:
        return int(ol.get(qn('w:val')))
    return None

# 收集 numPr（列表编号）
def numpr(p):
    pPr = p._p.find(qn('w:pPr'))
    if pPr is None:
        return None
    n = pPr.find(qn('w:numPr'))
    if n is None:
        return None
    ilvl = n.find(qn('w:ilvl'))
    numId = n.find(qn('w:numId'))
    return (
        int(ilvl.get(qn('w:val'))) if ilvl is not None else None,
        int(numId.get(qn('w:val'))) if numId is not None else None,
    )

# 斜体 run 检测
def italic_runs(p):
    out = []
    for r in p.runs:
        if r.text and r.font.italic:
            out.append(r.text)
    return out

# 字体（东亚）
def east_asia_font(p):
    fonts = set()
    for r in p.runs:
        rPr = r._r.find(qn('w:rPr'))
        if rPr is None:
            continue
        rf = rPr.find(qn('w:rFonts'))
        if rf is not None:
            ea = rf.get(qn('w:eastAsia'))
            if ea:
                fonts.add(ea)
    return fonts

print("\n========== 全部非空段落（带样式/大纲/编号/斜体） ==========")
n_heading = 0
list_para = 0
italic_headings = []
for i, p in enumerate(paras):
    t = p.text.strip()
    if not t:
        continue
    st = p.style.name if p.style else "?"
    ol = outline_level(p)
    np_ = numpr(p)
    tag = ""
    if st.startswith("Heading") or st.startswith("标题") or ol is not None:
        n_heading += 1
        tag = "[H]"
    if np_ is not None:
        list_para += 1
        tag += "[L%s/%s]" % np_
    ital = italic_runs(p)
    if ital and (st.startswith("Heading") or st.startswith("标题") or ol is not None):
        italic_headings.append((i, st, t, ital))
    ea = east_asia_font(p)
    extra = ""
    if ital:
        extra += " ITALIC=%r" % ("".join(ital)[:30],)
    if ea:
        extra += " FONT=%s" % (",".join(ea),)
    print("%4d %-16s ol=%s %s %s%s | %s" % (i, st, ol, tag, extra, "", t[:90]))

print("\n标题类段落数:", n_heading, " 带列表编号段落数:", list_para)
print("斜体标题数:", len(italic_headings))
for it in italic_headings[:20]:
    print("  斜体标题:", it)
