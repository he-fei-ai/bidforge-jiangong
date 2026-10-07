# -*- coding: utf-8 -*-
import re
from docx import Document

SRC = r"c:\Users\HEFEI\.trae-cn\attachments\6ac5f18c3e706c34f9ec0610\3597e0cd-ad7f-4049-8fc5-b5679b0ff62e_装饰装修专项施工方案_20261007_第4轮.docx"
doc = Document(SRC)

# 1. T10 全表（单位疑似错误）
print("=== T10 污染物浓度表全文 ===")
for r in doc.tables[10].rows:
    print(" | ".join(c.text.strip() for c in r.cells))

# 2. T21 医院表全文
print("\n=== T21 医院表全文 ===")
for r in doc.tables[21].rows:
    print(" | ".join(c.text.strip() for c in r.cells))

# 3. 所有表格单元格中的 Markdown 残留（**、- 开头、[ ]）
print("\n=== 表格单元格 Markdown 残留 ===")
cnt = 0
for ti, tbl in enumerate(doc.tables):
    for ri, row in enumerate(tbl.rows):
        for ci, cell in enumerate(row.cells):
            tx = cell.text
            if "**" in tx or re.search(r'(^|\n)\s*[-*]\s', tx) or "[" in tx and "]" in tx:
                print(f"T{ti} R{ri} C{ci}: {tx[:150]}")
                cnt += 1
                if cnt > 12: break
        if cnt > 12: break
    if cnt > 12: break

# 4. T24 全文（看残留结构）
print("\n=== T24 全文 ===")
for r in doc.tables[24].rows:
    print("名称:", r.cells[0].text.strip()[:60])
    print("内容:", r.cells[1].text.strip()[:300])
    print("---")

# 5. 正文中 0.07 / 0.05 出现位置
print("\n=== 甲醛 0.05/0.07 出现段落 ===")
for i, p in enumerate(doc.paragraphs):
    if ("0.05" in p.text or "0.07" in p.text) and ("甲醛" in p.text or "mg" in p.text):
        print(f"段{i}: {p.text.strip()[:150]}")

# 6. 工期天数出现位置
print("\n=== 工期天数 184/195/190/180/194 ===")
for i, p in enumerate(doc.paragraphs):
    m = re.search(r'(184|195|190|180|194)\s*(日历天|天)', p.text)
    if m:
        print(f"段{i}: {p.text.strip()[:120]}")

# 7. 检查章节内 H3 编号连续性（每个 H2 下 H3 是否从 1 开始）
print("\n=== 各 H2 下 H3 编号序列 ===")
cur_h2 = None
expect = 1
from docx.oxml.ns import qn
for p in doc.paragraphs:
    sn = p.style.name if p.style else ""
    t = p.text.strip()
    if sn == "Heading 2":
        cur_h2 = t
        expect = 1
        h2m = re.match(r'^(\d+)\b', t)
        print(f"\n[H2] {t}")
    elif sn == "Heading 3":
        m = re.match(r'^(\d+)\.(\d+)\b', t)
        if m:
            a, b = int(m.group(1)), int(m.group(2))
            flag = "" if b == expect else "  <-- 序号异常(期望 %d)" % expect
            print(f"    H3 {a}.{b} {t[len(m.group(0)):].strip()[:40]}{flag}")
            expect = b + 1
