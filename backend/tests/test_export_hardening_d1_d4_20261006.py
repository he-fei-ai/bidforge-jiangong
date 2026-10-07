"""导出 D1/D2/D3 加固护栏（R43 · 2026-10-06 收口轮）。

覆盖上一轮「遗留问题与后续建议」中已落地加固的三项：
  D1 · 跨章借图**零信号** → 计数 + 响应头 + 前端提示（不改变借用语义）
  D2 · 导出无进度 / 无取消 → 前端已落地，本文件锁住**服务端契约**部分
  D3 · 图表与配图两份平行实现 → 版式下沉为唯一实现
另含 D4 的「预检假阳性修复」的独立回归（与 parity 文件互补：
parity 文件锁两侧一致，本文件锁「默认配置下预检必须沉默」）。

设计原则（承前轮）：每条断言都指向**真实风险**，
且对"实现细节写错"与"意图写错"分别可判别。
"""
import ast
import inspect
import io
import os
import re

import pytest

from app.routers import export as E

REPO_BACKEND = os.path.dirname(os.path.dirname(os.path.dirname(
    os.path.abspath(E.__file__))))
REPO_ROOT = os.path.dirname(REPO_BACKEND)
EXPORT_PY = os.path.join(REPO_BACKEND, "app", "routers", "export.py")
FRONT_EXPORT_RESPONSE = os.path.join(
    REPO_ROOT, "frontend", "src", "utils", "exportResponse.ts")


def _sec(sid, pid, level, order, title, content=""):
    return {"id": sid, "parent_id": pid, "level": level,
            "sort_order": order, "title": title, "content": content}


# 三章均含 [CHART_TYPE: labor]；只有 secC 有登记 → secA/secB 走兜底借用
THREE_CHAPTER_SAME_TYPE = [
    _sec("secA", "", 1, 0, "第一章 劳动力组织", "## 总述\n[CHART_TYPE: labor]"),
    _sec("secB", "", 1, 1, "第二章 劳动力保障", "## 概述\n[CHART_TYPE: labor]"),
    _sec("secC", "", 1, 2, "第三章 劳动力培训", "## 安排\n[CHART_TYPE: labor]"),
]

CODE = "graph TD\n  A[劳动力计划] --> B[进场安排]\n"


def _png_bytes():
    from io import BytesIO

    from PIL import Image
    bio = BytesIO()
    Image.new("RGB", (200, 100), (200, 30, 30)).save(bio, "PNG")
    return bio.getvalue()


def _render(tmp_path, sections, chart_lookup, rendered_bytes, name="d1.docx"):
    import io as _io
    blocks = {s["id"]: E._parse_content_blocks(s.get("content") or "")
              for s in sections}
    cm = {}
    for s in sections:
        cm.setdefault(s.get("parent_id") or "", []).append(s)
    ids = {s["id"] for s in sections}
    prep = {
        "scheme": {"name": "S", "project_id": "p"},
        "roots": [s for s in sections
                  if not s.get("parent_id") or s["parent_id"] not in ids],
        "children_map": cm, "chart_lookup": chart_lookup,
        "rendered_bytes": rendered_bytes, "blocks_cache": blocks,
        "heading_styles": E._load_heading_styles({}), "image_bytes": {},
        "global_facts": [],
        "docx_options": {"font_name": "宋体", "font_size": 12.0,
                         "page_header": "", "page_footer": "",
                         "show_page_number": True, "show_title_page": False,
                         "show_toc": False, "bidder_name": "",
                         "page_break_before_chapter": True,
                         "line_spacing": 1.15, "page_number_style": "simple",
                         "toc_depth": 3},
    }
    out = str(tmp_path / name)
    stats = E._build_docx_sync(*E._build_docx_task(out, prep))
    return out, stats


@pytest.fixture(autouse=True)
def _reset_fix_stats():
    st = E._fix_stats()
    st.update({k: 0 for k in E._FIX_STATS_KEYS})
    yield
    st.update({k: 0 for k in E._FIX_STATS_KEYS})


# =========================================================================
# D1 · 跨章借图可观测
# =========================================================================
class TestCrossSectionBorrowIsObservable:
    def test_fix_stats_reports_borrowed_count(self, tmp_path):
        _out, stats = _render(
            tmp_path, THREE_CHAPTER_SAME_TYPE,
            {("secC", "labor"): CODE},
            {("labor", CODE): __import__("io").BytesIO(_png_bytes())})
        n = int((stats or {}).get("chart_fallback_borrowed") or 0)
        assert n == 2, f"secA/secB 应各记一次借图，实际 {n}"
        details = (stats or {}).get("chart_fallback_borrowed_details") or []
        assert len(details) == 2
        assert all(d["borrowed_from"] == "secC" for d in details)
        assert {d["section_id"] for d in details} == {"secA", "secB"}

    def test_self_registered_chapter_is_not_counted(self, tmp_path):
        """自有登记的章节不算借图（不得把正常图表误报成借图）。"""
        _out, stats = _render(
            tmp_path, THREE_CHAPTER_SAME_TYPE,
            {("secA", "labor"): CODE, ("secB", "labor"): CODE,
             ("secC", "labor"): CODE},
            {("labor", CODE): __import__("io").BytesIO(_png_bytes())})
        assert int((stats or {}).get("chart_fallback_borrowed") or 0) == 0

    def test_no_chart_means_no_borrow(self, tmp_path):
        tree = [_sec("a", "", 1, 0, "第一章", "## 总述\n[CHART_TYPE: labor]")]
        _out, stats = _render(tmp_path, tree, {}, {})
        assert int((stats or {}).get("chart_fallback_borrowed") or 0) == 0

    def test_merge_helper_is_single_source(self):
        """DOCX / PDF 两条链路的合流必须走同一个 helper（不得各写一遍）。

        ⚠️ 判据必须是**调用次数**而非"字符串出现过"：DOCX 分支有两处会回传
        渲染统计的响应头（坏缓存降级分支 + 正常 miss 分支），只查"出现过"
        会被另一处调用顶住 —— A/B 反向验证实测正是如此（摘掉 miss 分支的
        合流调用后，用例仍然绿）。
        """
        docx_src = inspect.getsource(E.export_docx)
        pdf_src = inspect.getsource(E.export_pdf)
        assert docx_src.count("_merge_borrowed_stats(render_stats, fix_stats)") >= 2, \
            "DOCX 的每个回传渲染统计的分支（降级 + miss）都必须合流借图计数"
        assert pdf_src.count("_merge_borrowed_stats(render_stats, fix_stats)") >= 1, \
            "PDF 链路未合流借图计数"
        # 三处合流点：DOCX 降级、DOCX miss、PDF
        whole = inspect.getsource(E)
        assert whole.count("_merge_borrowed_stats(render_stats, fix_stats)") == 3

    def test_merge_helper_shapes_the_payload(self):
        rs = {"fe": 1, "backend_ok": 2, "failed": 0}
        out = E._merge_borrowed_stats(rs, {
            "chart_fallback_borrowed": 3,
            "chart_fallback_borrowed_details": [{"section_id": "a"}],
        })
        assert out["fallback_borrowed"] == 3
        assert out["fallback_borrowed_details"][0]["section_id"] == "a"
        assert out["fe"] == 1 and out["backend_ok"] == 2, "既有渲染统计键被改动"
        # 无借图时也必须把键补 0（前端解析不必判 undefined）
        rs2 = {"fe": 0}
        assert E._merge_borrowed_stats(rs2, {})["fallback_borrowed"] == 0

    def test_fallback_choice_is_deterministic_by_document_order(self):
        """借图选择必须按**文档顺序**取最近的前一张，而非 DB 行插入顺序。"""
        # secC 登记了 codeL；secA/secB 都无登记 → 两者都应借 secC 的
        index = E._build_chart_type_index(
            {("secC", "labor"): CODE},
            order_rank={"secA": 0, "secB": 1, "secC": 2})
        code, src = E._find_fallback_code(index, "labor", "secA",
                                          {"secA": 0, "secB": 1, "secC": 2})
        assert (code, src) == (CODE, "secC")

    def test_fallback_prefers_nearest_preceding(self):
        """两章都有登记时，当前章应借**前面最近**的那张（而非全局第一张）。

        注意：真实流程里当前章**没有**自己的登记才会走兜底
        （`_resolve_chart_code` 先查 `chart_lookup[(sec_id, type)]`），
        故夹具里不含当前章。
        """
        other = "graph TD\n  C --> D\n"
        rank = {"s0": 0, "s1": 1, "s2": 2}
        # 故意让「全局第一张」属于 s0（顺序最靠前但离 s2 最远）
        index = E._build_chart_type_index(
            {("s0", "flowchart"): CODE, ("s1", "flowchart"): other}, rank)
        code, src = E._find_fallback_code(index, "flowchart", "s2", rank)
        assert src == "s1", f"应借文档顺序上最近的前一张，实际 {src}"
        assert code == other

    def test_fallback_falls_back_to_first_when_nothing_precedes(self):
        """候选全在自己之后时取第一张（保持"有图总比没图好"的兜底语义）。"""
        other = "graph TD\n  C --> D\n"
        rank = {"s1": 0, "s2": 1}
        index = E._build_chart_type_index(
            {("s2", "flowchart"): other}, rank)
        code, src = E._find_fallback_code(index, "flowchart", "s1", rank)
        assert (code, src) == (other, "s2")

    def test_fallback_missing_type_returns_empty_pair(self):
        assert E._find_fallback_code({}, "flowchart", "s1", {"s1": 0}) == ("", "")

    def test_resolver_is_single_source(self):
        """取码逻辑必须收敛到 `_resolve_chart_code`（两处各写一遍必然漂移）。"""
        prepare = inspect.getsource(E._prepare_export)
        render = inspect.getsource(E._build_docx_sync)
        assert "_resolve_chart_code(" in prepare
        assert "_resolve_chart_code(" in render
        # 旧的三元表达式不得复活
        assert "chart_lookup.get((sec[" not in prepare
        assert "chart_lookup.get((sec_id" not in render

    def test_index_carries_provenance(self):
        """倒排索引必须携带来源章节（否则无法上报借图）。"""
        idx = E._build_chart_type_index({("s9", "labor"): CODE})
        assert idx["labor"] == [("s9", CODE)]

    def test_index_dedups_same_section_same_code(self):
        idx = E._build_chart_type_index({("s9", "labor"): CODE,
                                         ("s9", "labor"): CODE})
        assert idx["labor"] == [("s9", CODE)]

    def test_index_skips_empty_codes(self):
        idx = E._build_chart_type_index({("s9", "labor"): ""})
        assert "labor" not in idx

    def test_frontend_surfaces_borrow_count(self):
        src = io.open(FRONT_EXPORT_RESPONSE, encoding="utf-8").read()
        assert "fallback_borrowed?:" in src, "前端 ChartRenderStats 缺 fallback_borrowed"
        assert "fallbackBorrowWarning" in src, "前端未提供借图告警文案"
        page = io.open(os.path.join(
            REPO_ROOT, "frontend", "src", "pages", "SchemeWorkbenchPage.tsx"),
            encoding="utf-8").read()
        assert "fallbackBorrowWarning(stats)" in page, "导出链路未消费借图告警"


# =========================================================================
# D3 · 版式实现唯一化
# =========================================================================
class TestImageInsertSingleImplementation:
    def test_shared_core_exists(self):
        assert callable(E._insert_image_with_caption)
        assert callable(E._write_chart_placeholder_line)

    def test_both_entries_delegate_to_core(self):
        for fn in (E._add_inline_chart_from_bytes,
                   E._add_illustration_from_bytes):
            assert "_insert_image_with_caption(" in inspect.getsource(fn), \
                f"{fn.__name__} 未下沉到唯一实现"

    def test_add_picture_appears_only_in_core(self):
        """⚠️ 防漂移核心断言：`doc.add_picture` 全文只允许出现在
        `_insert_image_with_caption` 一处。任何第二处都意味着版式修复又要改两遍
        （v20 的 DPI 修复历史上就改过两次）。

        判据用 **AST**（Attribute 调用）而非文本匹配 —— 纯文本扫描会被
        `_DocxBodyRollback` docstring 里那句 ``doc.add_picture()`` 命中而恒假失败
        （§5.14「护栏判据锚点选错，比不写护栏更糟」的又一例）。
        """
        tree = ast.parse(io.open(EXPORT_PY, encoding="utf-8").read())
        hits = []
        for node in ast.walk(tree):
            if (isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute)
                    and node.func.attr == "add_picture"):
                hits.append(node.lineno)
        assert len(hits) == 1, (
            f"doc.add_picture 调用点 {len(hits)} 处（行 {hits}），"
            "应唯一落在 _insert_image_with_caption 内")

    def test_add_picture_call_is_inside_core_function(self):
        """上一例的加强：确认唯一调用点确实位于核心实现体内。"""
        tree = ast.parse(io.open(EXPORT_PY, encoding="utf-8").read())
        for node in ast.walk(tree):
            if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)) \
                    and node.name == "_insert_image_with_caption":
                for sub in ast.walk(node):
                    if (isinstance(sub, ast.Call)
                            and isinstance(sub.func, ast.Attribute)
                            and sub.func.attr == "add_picture"):
                        return
        pytest.fail("核心实现内未找到 add_picture 调用")

    def test_fit_image_appears_only_in_core(self):
        src = io.open(EXPORT_PY, encoding="utf-8").read()
        hits = [i + 1 for i, ln in enumerate(src.split("\n"))
                if "_fit_image_cm(" in ln]
        # 1 处定义 + 1 处调用（定义行含 "def "）
        assert len(hits) == 2, f"_fit_image_cm 引用点异常：{hits}"

    def test_core_is_atomic_on_failure(self, monkeypatch):
        from docx import Document

        d = Document()
        orig = E._set_run_font

        def boom(run, *a, **kw):
            if (getattr(run, "text", "") or "").startswith("图 "):
                raise RuntimeError("simulated caption failure")
            return orig(run, *a, **kw)

        monkeypatch.setattr(E, "_set_run_font", boom)
        assert E._insert_image_with_caption(
            d, __import__("io").BytesIO(_png_bytes()), "图 1-1 X",
            log_label="probe") is False
        assert len(d.inline_shapes) == 0

    def test_chart_placeholder_text_unchanged(self):
        """占位红字形态是排查工具，不得被重构悄悄改掉。"""
        src = inspect.getsource(E._write_chart_placeholder_line)
        assert 'f"[{caption_text} — {reason}]"' in src

    def test_illustration_has_no_red_placeholder_by_default(self):
        """配图失败默认不写红字（交付文档里报错文本比少一张图更糟）。"""
        from docx import Document

        d = Document()
        ok = E._add_illustration_from_bytes(
            d, __import__("io").BytesIO(b"not-an-image"), "1-1", "配图")
        assert ok is False
        assert not [p for p in d.paragraphs if "[" in (p.text or "")]


# =========================================================================
# D4 补充 · 默认配置下预检必须沉默（假阳性修复的独立回归）
# =========================================================================
class TestDetectorFalsePositiveRegression:
    """与 parity 文件互补：这里只锁「默认配置下预检不得报不存在的重复」。"""

    def test_default_config_detector_is_silent(self):
        tree = [
            _sec("ch1", "", 1, 0, "第一章 工程概况",
                 "## 项目概况\n甲。\n\n## 建筑概况\n乙。\n"),
            _sec("a", "ch1", 2, 1, "工程规模", "## 规模描述\n丙。\n"),
            _sec("b", "ch1", 2, 2, "结构选型", ""),
            _sec("ch2", "", 1, 3, "第二章 施工计划", ""),
        ]
        from app.config import settings
        assert settings.body_subheading_demote_with_children is True, (
            "默认应为 True；若测试环境改了默认值，本用例的前提需复核")
        assert E._detect_duplicate_sections(tree) == []

    def test_detector_forwards_has_children(self):
        src = inspect.getsource(E._detect_duplicate_sections)
        assert "has_children=has_children_demoted" in src


# =========================================================================
# D2 · 服务端契约侧：进度/取消不依赖任何服务端改动，但导出必须可被中断
# =========================================================================
class TestExportCancellationContract:
    def test_export_routes_accept_abort_via_request_disconnect(self):
        """取消走客户端断连；两条路由都不得把断连当成业务成功写缓存。"""
        for fn in (E.export_docx, E.export_pdf):
            src = inspect.getsource(fn)
            assert "_prune_export_cache" in src
            # 缓存写入在构建完成之后（构建失败即无缓存行）
            assert src.index("_build_docx_sync") < src.index(
                "INSERT OR IGNORE INTO export_cache")

    def test_tmp_file_cleanup_on_build_failure(self):
        """构建异常必须清理临时文件（取消/断连时进程被杀也留 tmp，属已知取舍；
        正常异常路径必须清理）。"""
        for fn in (E.export_docx,):
            src = inspect.getsource(fn)
            assert "tmp_out_path.unlink(missing_ok=True)" in src

    @staticmethod
    def _page() -> str:
        return io.open(os.path.join(
            REPO_ROOT, "frontend", "src", "pages", "SchemeWorkbenchPage.tsx"),
            encoding="utf-8").read()

    def test_frontend_has_cancel_entry(self):
        """取消入口必须**真的调用 abort**。

        ⚠️ A/B 反向验证实测：只断言「函数名存在 / abort 字符串存在 / 按钮文案存在」
        会全部漏过 —— `handleExport` 自身也有一行 `exportAbortRef.current?.abort()`
        （用于覆盖上一次导出），文案常量也在 JSX 里，三者互相顶住。
        故判据必须锚定**函数体本体**。
        """
        page = self._page()
        assert "handleCancelExport" in page, "导出缺少取消入口"
        i = page.index("const handleCancelExport")
        body = page[i:page.index("\n  };", i)]
        assert "exportAbortRef.current?.abort()" in body, (
            "handleCancelExport 未真正触发 abort（函数存在但空实现）")
        assert "取消导出" in page, "取消按钮文案缺失"

    def test_frontend_has_progress_bar(self):
        page = self._page()
        assert "exportPct" in page, "导出缺少进度百分比状态"
        assert "<Progress" in page, "导出缺少进度条组件"

    def test_progress_resets_on_finish(self):
        """finally 必须把阶段与进度归零（否则下次导出残留上一次的进度条）。"""
        page = self._page()
        seg = page[page.index("} finally {"):page.index("const handleExportCheck")]
        assert 'setExportStage("idle")' in seg
        assert "setExportPct(0)" in seg

    def test_stage_state_has_real_consumer(self):
        """`exportStage` 必须被真正读取（防"声明了 state 却没人用"的死状态债）。"""
        page = self._page()
        assert re.search(r"EXPORT_STAGE_LABEL\[\s*exportStage\s*\]", page), \
            "exportStage 未被消费"
        assert re.search(r"EXPORT_STAGE_ORDER\[\s*exportStage\s*\]", page)