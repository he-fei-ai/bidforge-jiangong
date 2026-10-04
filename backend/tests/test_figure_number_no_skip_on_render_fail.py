"""回归测试：图号虚跳的两个独立场景（导出层图号分配，与正文生成核心无关）。

场景一 · 渲染失败图表不得占号（决策必须在占用图号之前）：
  旧实现先占号后插图，WEBP/AVIF/HEIC 这类"PIL 能解码、docx 插不进"的图，
  图号已被 +1 而 doc.add_picture 抛 UnrecognizedImageError 被吞 →
  「图号已占用但成稿无图」的静默虚跳。修复后 _chart_ok 含格式白名单，
  _image_format_supported 在占号前拦截（export.py 3159/3166、3202/3208）。

场景二 · 删除图表图号应紧凑重排（不出现断号）+ 指纹随之变化使缓存失效：
  图号是导出时按章节即时递增、无持久化；删除图表块后重新导出即自然重排。
  且 _content_fingerprint 把 section content 纳入指纹，删图致指纹变化 →
  缓存 miss → 不会命中含旧图的坏产物。

两类用例均为同步单测，不依赖数据库 / 渲染服务。
"""
import io
import os
import sys
import tempfile

import pytest
from docx import Document
from PIL import Image

sys.path.insert(0, r"J:\编程\专项方案工具箱\backend")

from app.routers.export import (  # noqa: E402
    _DEFAULT_HEADING_STYLES,
    _build_docx_sync,
    _content_fingerprint,
    _parse_content_blocks,
)


def _png_bytes() -> bytes:
    """永远可用的 docx 安全格式（PNG）。"""
    buf = io.BytesIO()
    Image.new("RGB", (60, 40), (30, 120, 200)).save(buf, "PNG")
    return buf.getvalue()


def _ppm_bytes() -> bytes:
    """PIL 可解码（format='PPM'）但不在 _DOCX_SAFE_IMAGE_FORMATS 白名单 ——
    确定性复现 WEBP/AVIF/HEIC 类「能解不能插」路径（无需依赖 WEBP 编码器）。
    """
    buf = io.BytesIO()
    Image.new("RGB", (60, 40), (10, 10, 10)).save(buf, "PPM")
    return buf.getvalue()


def _build(rendered_map: dict, content: str,
           chart_lookup: dict | None = None, fail_ph: bool = False,
           image_bytes: dict | None = None) -> Document:
    secs = [{"id": "s1", "parent_id": "", "level": 1, "sort_order": 0,
             "title": "第一章 工艺", "content": content}]
    out = os.path.join(tempfile.mkdtemp(prefix="fig_no_skip_"), "out.docx")
    _build_docx_sync(
        out_path=out, scheme={"name": "t", "project_id": "P1"},
        roots=secs, children_map={"": secs}, chart_lookup=chart_lookup or {},
        rendered_bytes=rendered_map, blocks_cache={},
        font_name="宋体", font_size=12, page_header="", page_footer="",
        show_page_number=False, show_title_page=False, show_toc=False,
        bidder_name="", heading_styles=_DEFAULT_HEADING_STYLES,
        page_break_before_chapter=True, line_spacing=1.15,
        page_number_style="simple", toc_depth=3, margins=None,
        cover_info=None, image_bytes=image_bytes or {}, global_facts=[],
        chart_fail_placeholder=fail_ph,
    )
    return Document(out)


def _captions(doc: Document) -> list[str]:
    return [p.text.strip() for p in doc.paragraphs
            if p.text.strip().startswith("图 ")]


def _chart_codes(content: str) -> list[str]:
    return [b["code"] for b in _parse_content_blocks(content)
            if b["type"] == "chart"]


# ============================================================
# 场景一：渲染失败图表不得占号
# ============================================================
class TestRenderFailureDoesNotOccupyFigureNumber:
    """决策必须在占用图号之前（export.py 3159/3166、3202/3208）。"""

    def test_good_png_occupies_number(self):
        """正向对照：渲染成功的图应当正常占号并产生图题。"""
        code = "graph TD\n  A[开始] --> B[结束]\n"
        content = "## 总述\n```mermaid\n" + code + "\n```"
        rendered = {("flowchart", c): io.BytesIO(_png_bytes()) for c in _chart_codes(content)}
        doc = _build(rendered, content)
        caps = _captions(doc)
        assert len(caps) == 1, caps
        assert caps[0].startswith("图 1-1"), caps

    def test_unsafe_format_does_not_occupy_number(self):
        """PPM 能解但 docx 插不进 → 占号前被 _image_format_supported 拦截。
        同时验证跳过图表时其孤儿引导语被回收（不出现「如下图所示」悬空引用）。"""
        code = "graph TD\n  A[开始] --> B[结束]\n"
        content = "施工工艺流程如下图所示：\n```mermaid\n" + code + "\n```"
        rendered = {("flowchart", c): io.BytesIO(_ppm_bytes()) for c in _chart_codes(content)}
        doc = _build(rendered, content)
        assert _captions(doc) == [], "不安全格式的图不该产生图号/图题（虚跳）"
        full = "\n".join(p.text for p in doc.paragraphs)
        assert "渲染失败" not in full and "插入失败" not in full, "不该写红字占位"
        assert "如下图所示" not in full, "跳过图表应回收孤儿引导语"

    def test_render_none_does_not_occupy_number(self):
        """渲染结果 None（渲染失败）同样不得占号。"""
        code = "graph TD\n  A[开始] --> B[结束]\n"
        content = "## 总述\n```mermaid\n" + code + "\n```"
        codes = _chart_codes(content)
        rendered = {("flowchart", c): None for c in codes}
        doc = _build(rendered, content)
        assert _captions(doc) == [], "渲染失败的图不该产生图号"
        full = "\n".join(p.text for p in doc.paragraphs)
        assert "渲染失败" not in full and "插入失败" not in full

    def test_fail_placeholder_true_shows_red_not_silent(self):
        """chart_fail_placeholder=True 恢复红字占位：图号被占用且红字含图号，
        属「可见错误」而非静默虚跳（与默认跳过形态互补）。"""
        code = "graph TD\n  A[开始] --> B[结束]\n"
        content = "## 总述\n```mermaid\n" + code + "\n```"
        rendered = {("flowchart", c): io.BytesIO(_ppm_bytes()) for c in _chart_codes(content)}
        doc = _build(rendered, content, fail_ph=True)
        full = "\n".join(p.text for p in doc.paragraphs)
        assert "插入失败" in full or "渲染失败" in full, "占位模式应显示红字报错"


# ============================================================
# 场景二：删除图表图号紧凑重排 + 指纹失效
# ============================================================
class TestDeletedChartFigureNumberCompaction:
    """删除图表后图号应紧凑重排（不出现断号）；指纹变化使缓存失效。"""

    @staticmethod
    def _codes(n: int) -> list[str]:
        return [f"graph TD\n  N{i}[节点{i}] --> M{i}[末端{i}]\n" for i in range(n)]

    @staticmethod
    def _content_with(codes: list[str]) -> str:
        mermaid = "\n".join("```mermaid\n" + c + "\n```" for c in codes)
        return "## 总述\n" + mermaid

    @staticmethod
    def _rendered_for(content: str) -> dict:
        return {("flowchart", c): io.BytesIO(_png_bytes())
                for c in _chart_codes(content)}

    def test_figure_numbers_compact_after_middle_deletion(self):
        codes3 = self._codes(3)
        content3 = self._content_with(codes3)
        doc3 = _build(self._rendered_for(content3), content3)
        caps3 = _captions(doc3)
        assert len(caps3) == 3, caps3
        assert caps3[0].startswith("图 1-1"), caps3
        assert caps3[1].startswith("图 1-2"), caps3
        assert caps3[2].startswith("图 1-3"), caps3

        # 删除中间一张（索引 1）
        codes2 = [codes3[0], codes3[2]]
        content2 = self._content_with(codes2)
        doc2 = _build(self._rendered_for(content2), content2)
        caps2 = _captions(doc2)
        # 关键：剩余两张应紧凑重排为 1-1 / 1-2，不得出现断号 1-3
        assert len(caps2) == 2, caps2
        assert caps2[0].startswith("图 1-1"), caps2
        assert caps2[1].startswith("图 1-2"), caps2
        assert not any("图 1-3" in c for c in caps2), "删除中间图后不得残留断号"

    def test_content_fingerprint_changes_on_chart_deletion(self):
        """删除图表后内容指纹必须变化，否则会命中含旧图的坏缓存。"""
        codes3 = self._codes(3)
        content3 = self._content_with(codes3)
        content2 = self._content_with([codes3[0], codes3[2]])

        sec3 = [{"id": "s1", "parent_id": "", "level": 1, "sort_order": 0,
                 "title": "第一章 工艺", "content": content3}]
        sec2 = [dict(sec3[0], content=content2)]

        def _fp(sections, codes):
            return _content_fingerprint({
                "sections": sections,
                "chart_fp": [("s1", "flowchart", c) for c in codes],
                "config": {}, "fe_codes": [], "global_facts": [],
            })[1]

        h3 = _fp(sec3, codes3)
        h2 = _fp(sec2, [codes3[0], codes3[2]])
        assert h3 != h2, "删除图表后内容指纹必须变化，否则会命中含旧图的坏缓存"


# ============================================================
# 可修复项（2026-10-03）：图字节格式合法但文件损坏 → doc.add_picture 抛异常
# 旧实现图号已 +1 却无图（错号/虚跳）。修复后插入失败回退图号、默认模式不写红字。
# 用 monkeypatch 让 doc.add_picture 抛异常，确定性模拟「格式合法但文件损坏」路径，
# 无需构造真实损坏图片。
# ============================================================
class TestCorruptImageRollback:
    """插入失败（格式合法但文件损坏）必须回退图号，不得「占号却无图」。"""

    @staticmethod
    def _boom(self, *args, **kwargs):
        raise Exception("simulated insert failure")

    def test_corrupt_chart_default_mode_no_number(self, monkeypatch):
        import docx.document as _docx_document
        monkeypatch.setattr(_docx_document.Document, "add_picture", self._boom)
        code = "graph TD\n  A[开始] --> B[结束]\n"
        content = "施工工艺流程如下图所示：\n```mermaid\n" + code + "\n```"
        rendered = {("flowchart", c): io.BytesIO(_png_bytes())
                    for c in _chart_codes(content)}
        doc = _build(rendered, content)
        assert _captions(doc) == [], "损坏图表不得占号（回退后无图题）"
        full = "\n".join(p.text for p in doc.paragraphs)
        assert "插入失败" not in full and "渲染失败" not in full, "默认模式不写红字"
        assert "如下图所示" not in full, "损坏图表应回收孤儿引导语"

    def test_corrupt_chart_placeholder_mode_shows_red(self, monkeypatch):
        import docx.document as _docx_document
        monkeypatch.setattr(_docx_document.Document, "add_picture", self._boom)
        code = "graph TD\n  A[开始] --> B[结束]\n"
        content = "## 总述\n```mermaid\n" + code + "\n```"
        rendered = {("flowchart", c): io.BytesIO(_png_bytes())
                    for c in _chart_codes(content)}
        doc = _build(rendered, content, fail_ph=True)
        full = "\n".join(p.text for p in doc.paragraphs)
        assert "插入失败" in full or "渲染失败" in full, \
            "占位模式应显示红字（图号被占用、可见错误）"

    def test_corrupt_ai_image_default_mode_no_number(self, monkeypatch):
        import docx.document as _docx_document
        monkeypatch.setattr(_docx_document.Document, "add_picture", self._boom)
        url = "https://example.com/fig.png"
        content = "现场示意图如下图所示：\n![示意图](%s)\n" % url
        # 合法 PNG 使 _chart_ok 通过，但插入时抛异常（损坏场景）
        doc = _build({}, content, image_bytes={url: io.BytesIO(_png_bytes())})
        assert _captions(doc) == [], "损坏配图不得占号（回退后无图题）"
        full = "\n".join(p.text for p in doc.paragraphs)
        assert "插入失败" not in full and "渲染失败" not in full, "默认模式不写红字"
        assert "如下图所示" not in full, "损坏配图应回收孤儿引导语"
