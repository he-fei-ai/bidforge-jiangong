"""export_docx 端到端冒烟测试

目标：确认导出主流程在真实 DB + 章节树 / 图表 / 子标题嵌套场景下
不会抛异常，并产出 docx 文件。

- 仅校验流程不崩溃、返回 ok、文件存在、关键编号正确；
- 图表渲染依赖外部服务，测试环境可能不可用 —— 渲染失败时导出会插入
  “渲染失败”提示而非崩溃，因此不校验图片像素。

⚠️ 关于「PIL 兜底 / 渲染失败产物形态」的契约：
  · `allow_pil_fallback` 默认 False（与前端导出表单 Switch 初值、tooltip 一致），
    Mermaid 类图表在 HTTP Service 不可用时**不静默 PIL 自绘**；
    `test_export_docx_smoke` 显式传 True，守护「离线确实能出图」路径。
  · ✅ 2026-09-19（v12）契约变更：渲染失败的图表**默认整块跳过且不占图号**，
    交付文档不再出现红字「渲染失败」段（评审场景报错文字比缺图更糟）；
    "不静默"承诺转移到系统侧信号（X-Chart-Render-Stats / X-Content-Audit /
    预检 / 日志）。传 `chart_fail_placeholder=True` 恢复占位形态，见
    `test_export_docx_optin_placeholder_keeps_v70_contract`。
"""
import os

import pytest
from app.routers.export import export_docx


async def _seed_export_fixture(db_conn):
    """写入冒烟测试用的方案 / 章节树 / 图表预测。"""
    # 方案
    await db_conn.execute(
        "INSERT INTO schemes (id, project_id, name) VALUES (?,?,?)",
        ("s1", "p1", "测试专项方案"),
    )
    # 章节树：第一章 -> 1.1 工程概况；第一章正文含 Markdown 子标题与图表标记
    await db_conn.execute(
        "INSERT INTO sections (id, scheme_id, parent_id, sort_order, level, title, status, content) "
        "VALUES (?,?,?,?,?,?,?,?)",
        ("c1", "s1", "", 0, 1, "第一章 总体概述", "generated",
         "## 总体安排\n这里是第一章正文。\n[CHART_TYPE: flowchart]\n"),
    )
    await db_conn.execute(
        "INSERT INTO sections (id, scheme_id, parent_id, sort_order, level, title, status, content) "
        "VALUES (?,?,?,?,?,?,?,?)",
        ("c1-1", "s1", "c1", 1, 2, "1.1 工程概况", "generated",
         "### 关键指标\n工程概况内容。\n"),
    )
    # 图表预测（与 c1 的 [CHART_TYPE: flowchart] 标记对应）
    await db_conn.execute(
        "INSERT INTO chart_predictions (id, section_id, scheme_id, chart_type, needed, purpose, priority, status, data_json) "
        "VALUES (?,?,?,?,?,?,?,?,?)",
        ("cp1", "c1", "s1", "flowchart", 1, "施工流程图", 3, "generated",
         '{"mermaid_code":"graph TD\\n A[开始]-->B[结束]"}'),
    )
    await db_conn.commit()


async def _run_export(db_conn, config: dict):
    await _seed_export_fixture(db_conn)
    body = {
        "project_name": "示例项目",
        "bid_section": "第一标段",
        "config": config,
        "chart_images": [],
        "fe_codes": [],
    }
    return await export_docx("s1", body, db=db_conn)


@pytest.mark.asyncio
async def test_export_docx_smoke(db_conn):
    # ✅ 显式开启 PIL 兜底：本用例守护「内置渲染器离线出图」这条路径
    resp = await _run_export(db_conn, {"allow_pil_fallback": True})

    # 成功路径返回 FileResponse（路径在 .path 上）
    from fastapi.responses import FileResponse
    assert isinstance(resp, FileResponse)
    path = resp.path
    assert path and os.path.exists(path)

    # 读取生成的 docx，校验嵌套子标题编号已生效
    from docx import Document
    doc = Document(path)
    para = [(p.style.name, p.text) for p in doc.paragraphs]
    joined = "\n".join(t for _, t in para)

    # ✅ E3（2026-09-25）：第一章有 DB 子章节 → 正文子标题降级为节内 body 命名空间
    # （1）、2）、…，与 DB 子章节的 X.X 命名空间彻底隔离——不再撞号。
    assert "1）、总体安排" in joined, f"未找到降级后的正文子标题编号，实际内容:\n{joined}"
    # c1-1 是叶子章节（无 DB 子章节），正文子标题不降级 → 保持旧点分格式
    assert "1.1 关键指标" in joined, f"未找到二级嵌套子标题，实际内容:\n{joined}"
    # 图号应生成为 "1-1"（章节号-序号）
    assert "1-1" in joined, f"未找到图号，实际内容:\n{joined}"
    # ✅ allow_pil_fallback=True 时应嵌入真实图片而非「渲染失败」占位
    assert "渲染失败" not in joined, f"图表未成功渲染（出现失败占位），实际内容:\n{joined}"
    import zipfile as _zf
    with _zf.ZipFile(path) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
    assert media, "未找到嵌入的图表图片"
    # ✅ 图题必须是「图 X-Y 图名」顺序（旧实现拼成"图名 1-1"）
    assert "图 1-1" in joined, f"图题格式不正确，实际内容:\n{joined}"


@pytest.mark.asyncio
async def test_export_heading_border_enabled(db_conn):
    """✅ 2026-09-22 引入（对齐 OpenBidKit 章框 heading_border）：
    开启时一级章节标题（Heading 1）应带底部边框 w:pBdr/w:bottom。"""
    from docx import Document
    from docx.oxml.ns import qn
    from fastapi.responses import FileResponse

    resp = await _run_export(db_conn, {"heading_border": True})
    assert isinstance(resp, FileResponse)
    assert resp.path and os.path.exists(resp.path)

    doc = Document(resp.path)
    found = False
    bordered = False
    for p in doc.paragraphs:
        if p.style is not None and p.style.name == "Heading 1":
            found = True
            pPr = p._p.find(qn("w:pPr"))
            pbd = pPr.find(qn("w:pBdr")) if pPr is not None else None
            bordered = pbd is not None and pbd.find(qn("w:bottom")) is not None
            break
    assert found, "未找到一级章节标题（Heading 1）"
    assert bordered, "开启 heading_border 后一级标题应带底部边框"


@pytest.mark.asyncio
async def test_export_heading_border_disabled_by_default(db_conn):
    """默认关闭 heading_border（向后兼容）：一级章节标题不应出现边框。"""
    from docx import Document
    from docx.oxml.ns import qn
    from fastapi.responses import FileResponse

    resp = await _run_export(db_conn, {})
    assert isinstance(resp, FileResponse)
    assert resp.path and os.path.exists(resp.path)

    doc = Document(resp.path)
    for p in doc.paragraphs:
        if p.style is not None and p.style.name == "Heading 1":
            pPr = p._p.find(qn("w:pPr"))
            pbd = pPr.find(qn("w:pBdr")) if pPr is not None else None
            assert pbd is None or pbd.find(qn("w:bottom")) is None, \
                "默认不应出现章框边框（保持旧行为）"
            break


@pytest.mark.asyncio
async def test_export_docx_default_no_pil_fallback_skips_failed_chart(db_conn):
    """✅ 契约变更（2026-09-19，v12）：默认（不降级）策略下渲染失败的图表被**跳过
    且不占图号**，交付文档不出现红字「渲染失败」段（实测取证：红字报错段在评审
    场景比缺图更糟）。导出本身不崩溃，失败信息仍通过响应头 / 日志 / 预检外显。
    """
    resp = await _run_export(db_conn, {})

    from fastapi.responses import FileResponse
    assert isinstance(resp, FileResponse)
    assert resp.path and os.path.exists(resp.path)

    from docx import Document
    joined = "\n".join(p.text for p in Document(resp.path).paragraphs)
    # 文档结构必须完整（编号照常生成），但不再有红字占位段
    # ✅ E3：第一章有 DB 子章节 → 正文子标题降级为 1）、2）、… 格式
    assert "1）、总体安排" in joined, joined

    import zipfile as _zf
    with _zf.ZipFile(resp.path) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]

    if media:
        # 本机有 Mermaid HTTP Service：正常出图
        assert "渲染失败" not in joined, joined
    else:
        # 无服务：失败图被跳过 → 既无图也无占位文字，图号不虚跳
        assert "渲染失败" not in joined, f"交付文档不应出现红字占位:\n{joined}"
        assert "图 1-1" not in joined, f"跳过的图不应占用图号:\n{joined}"
        # 系统侧信号必须保留（"不静默"承诺转移到响应头）
        import json as _json
        from urllib.parse import unquote
        stats = _json.loads(resp.headers["X-Chart-Render-Stats"])
        assert stats["failed"] >= 1, stats


@pytest.mark.asyncio
async def test_export_docx_optin_placeholder_keeps_v70_contract(db_conn):
    """chart_fail_placeholder=True 恢复 V7.0 形态：渲染失败写红字占位（排查用）。"""
    resp = await _run_export(db_conn, {"chart_fail_placeholder": True})

    from docx import Document
    joined = "\n".join(p.text for p in Document(resp.path).paragraphs)

    import zipfile as _zf
    with _zf.ZipFile(resp.path) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]

    if not media:
        assert "图 1-1" in joined, f"占位模式下图号照常分配:\n{joined}"
        assert "渲染失败" in joined, f"占位模式应写入红字提示:\n{joined}"



# ===========================================================================
# 新版式能力（章节分页 / 域自动刷新 / 页脚域结果 / 代码块 / 表格表头 / 链接净化）
# ===========================================================================
import zipfile
from io import BytesIO

from app.routers.export import _build_docx_sync, _load_heading_styles, _parse_content_blocks
from docx import Document

_CHART_CODE = "graph TD\n A-->B"
_CHAPTER1_CONTENT = (
    "本章说明施工总体安排。\n"
    "\n"
    "## 施工顺序\n"
    "1. 参见 [附件1](http://example.com/a) 与 `PB-2026` 规范。\n"
    "\n"
    "> 雨季施工需加密监测频率。\n"
    "\n"
    "---\n"
    "\n"
    "| 阶段 | 工期 |\n"
    "|---|---|\n"
    "| 基础 | 30天 |\n"
    "| 主体 | **60**天 |\n"
    "\n"
    "```bash\n"
    "  npm run build -- --mode prod\n"
    "```\n"
    "\n"
    "```mermaid\n"
    "graph TD\n"
    " A-->B\n"
    "```\n"
)


def _png_bytes() -> bytes:
    from PIL import Image
    buf = BytesIO()
    Image.new("RGB", (60, 40), (200, 210, 230)).save(buf, format="PNG")
    return buf.getvalue()


def _build_layout_docx(tmp_path, *, show_title_page=True, show_toc=True,
                       chapter_page_break=True, page_number_style="page_of_total"):
    sections = [
        {"id": "c1", "parent_id": "", "sort_order": 0, "level": 1,
         "title": "第一章 总体概述", "content": _CHAPTER1_CONTENT},
        {"id": "c2", "parent_id": "", "sort_order": 1, "level": 1,
         "title": "第二章 施工部署", "content": "第二章正文。"},
    ]
    blocks = {s["id"]: _parse_content_blocks(s["content"]) for s in sections}
    out = str(tmp_path / "layout.docx")
    _build_docx_sync(
        out,
        {"id": "s1", "project_id": "p1", "name": "深基坑开挖专项施工方案"},
        sections, {},                                  # roots / children_map
        {},                                            # chart_lookup（图表走内联码）
        {("flowchart", _CHART_CODE): BytesIO(_png_bytes())},
        blocks,
        "微软雅黑", 12, "深基坑专项方案", "", True,
        show_title_page, show_toc,
        # ✅ 追加参数一律用关键字：_build_docx_sync 的 bidder_name 被插在
        #    heading_styles 之前（export.py 注释自己写了「追加参数统一放末尾，
        #    插入中间会错位」，但没遵守），本文件的位置参数因此整体右移一位
        #    —— line_spacing 收到 "simple" 字符串，Emu("simple" * 152400) 抛
        #    ValueError，本文件 17 例连锁失败。改用关键字后不再受影响。
        bidder_name="",
        heading_styles=_load_heading_styles({}),
        page_break_before_chapter=chapter_page_break,
        line_spacing=1.5,
        page_number_style=page_number_style,
        toc_depth=3,
    )
    return out


def _part(path, prefix):
    """读取 docx 中第一个匹配前缀的 XML 部件（如 word/footer*）"""
    with zipfile.ZipFile(path) as z:
        name = next(n for n in z.namelist() if n.startswith(prefix))
        return z.read(name).decode("utf-8")


def test_docx_layout_enhancements(tmp_path):
    """一次构建校验：章节分页 / 目录页码域 / 图题 / 代码块底纹 / 表格表头 / 链接净化"""
    path = _build_layout_docx(tmp_path)
    doc_xml = _part(path, "word/document.xml")
    settings_xml = _part(path, "word/settings.xml")
    footer_xml = _part(path, "word/footer")

    # 1) 图题规范化为「图 X-Y 图名」，且图片真正嵌入
    assert "图 1-1 施工流程图" in doc_xml
    assert "<w:drawing>" in doc_xml

    # 2) 一级章节另起一页：仅第二章分页（首章不分页，避免封面/目录后多出空白页）
    assert doc_xml.count("pageBreakBefore") == 1

    # 3) 封面 + 目录各有一个分页符
    assert doc_xml.count('w:type="page"') == 2
    # 4) 有封面 → 首页不显示页眉页脚（titlePg）
    assert "<w:titlePg" in doc_xml

    # 5) 声明"打开时更新域"，且必须排在 w:compat 之前（否则 Word 报文件损坏）
    assert "<w:updateFields" in settings_xml
    assert settings_xml.index("<w:updateFields") < settings_xml.index("<w:compat")

    # 6) 页脚：PAGE + NUMPAGES 域，且带 separate + 占位结果（旧实现缺 separate → 未刷新时页码空白）
    assert "PAGE" in footer_xml and "NUMPAGES" in footer_xml
    assert footer_xml.count('w:fldCharType="separate"') == 2

    # 7) 代码块：浅灰底纹 + 四边边框 + 等宽字体 + 行首缩进保留（NBSP）
    assert "F2F2F2" in doc_xml
    assert "Consolas" in doc_xml
    assert "\u00a0" in doc_xml

    # 8) 表格：表头底纹 + 跨页重复
    assert "B6B1D1" in doc_xml
    assert "<w:tblHeader" in doc_xml

    # 9) 行内 Markdown 净化：链接只留显示文本，代码内联用等宽字体
    assert "附件1" in doc_xml
    assert "http://example.com/a" not in doc_xml
    assert "[附件1]" not in doc_xml


def test_docx_page_break_can_be_disabled(tmp_path):
    """chapter_page_break=False 时不应插入章节分页"""
    path = _build_layout_docx(tmp_path, chapter_page_break=False)
    assert "pageBreakBefore" not in _part(path, "word/document.xml")


def test_docx_simple_page_number_style(tmp_path):
    """page_number_style=simple 时不使用 NUMPAGES 域"""
    path = _build_layout_docx(tmp_path, page_number_style="simple")
    footer_xml = _part(path, "word/footer")
    assert " PAGE " in footer_xml
    assert "NUMPAGES" not in footer_xml


def test_docx_without_title_page_no_titlepg(tmp_path):
    """没有封面时不设置首页不同，也不插入封面的分页符"""
    path = _build_layout_docx(tmp_path, show_title_page=False, show_toc=False)
    doc_xml = _part(path, "word/document.xml")
    assert "<w:titlePg" not in doc_xml
    assert doc_xml.count('w:type="page"') == 0


# ===========================================================================
# v7 新增能力验证：有序列表自动编号 / 表格列宽自适应 / 图片按比例不放大 /
# 封面项目信息表 / 页边距可配置
# ===========================================================================
def _build_single(tmp_path, content: str, **overrides) -> str:
    """构建单章节 DOCX，便于针对性验证某一项排版能力。"""
    sections = [{
        "id": "c1", "parent_id": "", "sort_order": 0, "level": 1,
        "title": "第一章 测试章节", "content": content,
    }]
    blocks = {"c1": _parse_content_blocks(content)}
    out = str(tmp_path / "single.docx")
    opts = dict(
        font_name="宋体", font_size=12, page_header="", page_footer="",
        show_page_number=False, show_title_page=False, show_toc=False,
        heading_styles=_load_heading_styles({}),
        page_break_before_chapter=True, line_spacing=1.5,
        page_number_style="simple", toc_depth=3,
    )
    opts.update(overrides)
    _build_docx_sync(
        out, {"id": "s1", "project_id": "p1", "name": "测试方案"},
        sections, {}, {}, {}, blocks,
        opts["font_name"], opts["font_size"], opts["page_header"], opts["page_footer"],
        opts["show_page_number"], opts["show_title_page"], opts["show_toc"],
        bidder_name="",
        heading_styles=opts["heading_styles"],
        page_break_before_chapter=opts["page_break_before_chapter"],
        line_spacing=opts["line_spacing"],
        page_number_style=opts["page_number_style"],
        toc_depth=opts["toc_depth"],
        margins=opts.get("margins"), cover_info=opts.get("cover_info"),
    )
    return out


def test_docx_hr_not_rendered_as_line(tmp_path):
    """回归（2026-09-20）：AI 正文中的 --- 分隔线不应画成横线。

    旧实现在 hr 分支画「紧凑空段落 + BFBFBF 下边框」，正式交付文档中出现
    突兀横线（如第 2 章导语与 2.1 小节之间）。解析仍识别 hr，渲染侧跳过。
    """
    content = "本章为技术标准与规范导语。\n\n---\n\n1.1 《施工脚手架通用规范》应严格执行。\n"
    out = _build_single(tmp_path, content)
    doc_xml = _part(out, "word/document.xml")
    # 横线专属颜色不再出现（hr 是唯一使用 BFBFBF 底边框的分支）
    assert "BFBFBF" not in doc_xml, "正文 --- 仍被渲染为横线"
    # 前后正文段落必须保留（跳过 hr 不吞内容）
    joined = "\n".join(p.text for p in Document(out).paragraphs)
    assert "本章为技术标准与规范导语。" in joined
    assert "1.1 《施工脚手架通用规范》应严格执行。" in joined


def test_docx_ordered_list_auto_renumber(tmp_path):
    """AI 输出重复/跳号的有序列表，导出应自动连续编号为 1. 2. 3. 4."""
    content = (
        "正文。\n"
        "1. 第一步内容\n"
        "1. 第二步内容（AI 误写重复编号）\n"
        "1. 第三步内容\n"
        "1. 第四步内容\n"
    )
    out = _build_single(tmp_path, content)
    paras = [p.text for p in Document(out).paragraphs]
    assert "1. 第一步内容" in paras, paras
    assert "2. 第二步内容（AI 误写重复编号）" in paras, paras
    assert "3. 第三步内容" in paras, paras
    assert "4. 第四步内容" in paras, paras


def test_docx_table_column_widths(tmp_path):
    """表格应声明固定布局 + 总宽（按内容分配列宽，不等宽挤压）。"""
    content = (
        "| 短 | 较长较长较长较长较长较长较长较长较长较长 |\n"
        "|---|---|\n"
        "| a | b |\n"
        "| c | d |\n"
    )
    out = _build_single(tmp_path, content)
    doc_xml = _part(out, "word/document.xml")
    assert "w:tblLayout" in doc_xml and 'w:type="fixed"' in doc_xml
    assert "w:tblW" in doc_xml


def test_docx_image_small_not_enlarged(tmp_path):
    """窄图（100px 宽）不应被放大到 16cm，而是保持原始比例尺寸。"""
    from PIL import Image as PILImage
    buf = BytesIO()
    PILImage.new("RGB", (100, 40), (200, 210, 230)).save(buf, format="PNG")
    small = buf.getvalue()
    content = "```mermaid\ngraph TD\n A-->B\n```"
    sections = [{
        "id": "c1", "parent_id": "", "sort_order": 0, "level": 1,
        "title": "第一章 图", "content": content,
    }]
    blocks = {"c1": _parse_content_blocks(content)}
    out = str(tmp_path / "img.docx")
    _build_docx_sync(
        out, {"id": "s1", "project_id": "p1", "name": "t"},
        sections, {}, {}, {("flowchart", "graph TD\n A-->B"): BytesIO(small)},
        blocks, "宋体", 12, "", "", False, False, False,
        bidder_name="",
        heading_styles=_load_heading_styles({}),
        page_break_before_chapter=True,
        line_spacing=1.5,
        page_number_style="simple",
        toc_depth=3,
        margins=None, cover_info=None,
    )
    doc_xml = _part(out, "word/document.xml")
    import re as _re
    m = _re.search(r'wp:extent cx="(\d+)"', doc_xml)
    assert m, "未找到图片 extent"
    cx = int(m.group(1))
    # 100px @96dpi ≈ 952500 EMU；16cm = 5760000 EMU。窄图应远小于 16cm。
    assert cx < 5760000, f"窄图被放大：cx={cx}"


def test_docx_cover_info_table(tmp_path):
    """开启封面页且提供 cover_info 时，文档应含项目信息表（标签/值文本）。"""
    out = _build_single(
        tmp_path, "正文。",
        show_title_page=True,
        cover_info={"工程名称": "示例工程", "编制单位": "示例建设公司", "方案编号": "FA-1"},
    )
    doc = Document(out)
    texts = [p.text for p in doc.paragraphs]
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                texts.append(cell.text)
    joined = "\n".join(texts)
    assert "示例工程" in joined
    assert "示例建设公司" in joined
    assert _part(out, "word/document.xml").count("<w:tbl>") >= 1  # 封面信息表


def test_docx_margins_applied(tmp_path):
    """传入页边距应写入 sectPr 的 pgMar（左 3.0cm ≈ 1701 twips）。"""
    out = _build_single(
        tmp_path, "正文。",
        margins={"left": 3.0, "right": 3.0, "top": 2.0, "bottom": 2.0})
    doc_xml = _part(out, "word/document.xml")
    assert "w:pgMar" in doc_xml
    assert "1701" in doc_xml  # 3.0cm * 567 twips/cm = 1701


# ===========================================================================
# 导出文件名规则：专项方案名称 + 日期 + 导出轮次
# ===========================================================================
def test_export_filename_truncation_keeps_round():
    """超长方案名截断后必须保留「_日期_第N轮」后缀（否则轮次信息丢失、同日重名）。"""
    from app.routers.export import _build_export_filename
    name = _build_export_filename("超长方案名" * 60, "docx", 12)
    assert name.endswith("_第12轮.docx"), name
    assert len(name) <= 120, f"文件名超长：{len(name)}"


@pytest.mark.asyncio
async def test_export_filename_rule_with_round(db_conn):
    """端到端契约：文件名 = 方案名称 + 导出日期 + 导出轮次，且每次导出轮次自增。

    轮次按「导出请求」计数（缓存命中也算一轮），保证同一天多次导出的文件名不重复。
    """
    from datetime import datetime
    from urllib.parse import unquote

    await _seed_export_fixture(db_conn)
    body = {"config": {"allow_pil_fallback": True}, "chart_images": [], "fe_codes": []}
    day = datetime.now().strftime("%Y%m%d")

    for expected in (1, 2):
        resp = await export_docx("s1", body, db=db_conn)
        assert resp.headers["X-Export-Round"] == str(expected)
        # 前端 blob 下载据此设置 a.download，故必须是可直接解码的文件名
        assert unquote(resp.headers["X-Export-Filename"]) == \
            f"测试专项方案_{day}_第{expected}轮.docx"
        # 后端 Content-Disposition 与之一致（非 ASCII 走 RFC 5987）
        assert f"第{expected}轮.docx" in unquote(resp.headers["content-disposition"])

    cur = await db_conn.execute("SELECT export_round FROM schemes WHERE id='s1'")
    assert (await cur.fetchone())[0] == 2


def test_docx_chart_caption_follows_font_config(tmp_path):
    """✅ 增强：图题字体/字号跟随导出配置（旧实现硬编码「宋体 10.5pt」，
    用户把正文设为微软雅黑/四号后图题会与正文明显不一致）。"""
    content = "```mermaid\ngraph TD\n A-->B\n```"
    sections = [{
        "id": "c1", "parent_id": "", "sort_order": 0, "level": 1,
        "title": "第一章 图", "content": content,
    }]
    blocks = {"c1": _parse_content_blocks(content)}
    out = str(tmp_path / "caption.docx")
    _build_docx_sync(
        out, {"id": "s1", "project_id": "p1", "name": "t"},
        sections, {}, {}, {("flowchart", _CHART_CODE): BytesIO(_png_bytes())},
        blocks, "微软雅黑", 14, "", "", False, False, False,
        bidder_name="",
        heading_styles=_load_heading_styles({}),
        page_break_before_chapter=True,
        line_spacing=1.5,
        page_number_style="simple",
        toc_depth=3,
        margins=None, cover_info=None,
    )
    from docx.oxml.ns import qn
    doc = Document(out)
    caps = [p for p in doc.paragraphs if p.text.startswith("图 1-1")]
    assert caps, [p.text for p in doc.paragraphs]
    run = caps[0].runs[0]
    assert run.font.name == "微软雅黑"
    assert run._element.rPr.rFonts.get(qn("w:eastAsia")) == "微软雅黑"
    assert run.font.size.pt == 14


# ===========================================================================
# v9 新增：中文枚举符样式保留 / 有序列表续号 / 表格题注
# ===========================================================================
def test_docx_ordered_list_continues_after_unordered_item(tmp_path):
    """✅ 优化：有序列表被无序子项打断后序号继续递增（旧实现从 1 重排）"""
    content = (
        "1. 第一步：测量放线\n"
        "- 注意事项：避开雨天\n"
        "2. 第二步：基坑开挖\n"
    )
    out = _build_single(tmp_path, content)
    paras = [p.text for p in Document(out).paragraphs]
    assert "1. 第一步：测量放线" in paras, paras
    assert "2. 第二步：基坑开挖" in paras, paras
    assert "• 注意事项：避开雨天" in paras, paras
    assert not any(t.startswith("1. 第二步") for t in paras), paras


def test_docx_chinese_list_marker_preserved(tmp_path):
    """✅ 优化：中文枚举符外观被保留（不再一律拍平成 "1. "），且重复编号自动修正"""
    content = (
        "（一）设计标准\n"
        "（二）材料要求\n"
        "\n"
        "正文过渡段落。\n"
        "\n"
        "1、施工准备\n"
        "1、测量放线\n"
        "\n"
        "正文过渡段落二。\n"
        "\n"
        "（1）土方开挖\n"
        "（1）基坑支护\n"
    )
    out = _build_single(tmp_path, content)
    paras = [p.text for p in Document(out).paragraphs]
    assert "（一）设计标准" in paras, paras
    assert "（二）材料要求" in paras, paras
    assert "1、施工准备" in paras, paras
    assert "2、测量放线" in paras, paras          # AI 重复编号被修正，顿号样式保留
    assert "（1）土方开挖" in paras, paras
    assert "（2）基坑支护" in paras, paras        # 全角括号样式保留


def test_docx_table_caption_numbered_above_table(tmp_path):
    """✅ 新增：表格题注「表 X-Y 表名」渲染在表格上方，原表名行不再重复出现"""
    content = (
        "表3-1 主要施工机械设备表\n"
        "\n"
        "| 序号 | 名称 |\n"
        "|---|---|\n"
        "| 1 | 挖掘机 |\n"
    )
    out = _build_single(tmp_path, content)
    texts = [p.text for p in Document(out).paragraphs]
    assert "表 1-1 主要施工机械设备表" in texts, texts
    assert not any(t.strip() == "表3-1 主要施工机械设备表" for t in texts), texts
    doc_xml = _part(out, "word/document.xml")
    assert doc_xml.index("表 1-1 主要施工机械设备表") < doc_xml.index("<w:tbl>")


def test_docx_table_caption_sequence_per_chapter(tmp_path):
    """同章内多张表格的题注序号应连续为 1-1 / 1-2"""
    content = (
        "表1 设备表\n\n| A | B |\n|---|---|\n| 1 | 2 |\n"
        "\n正文说明。\n\n"
        "表2 材料表\n\n| A | B |\n|---|---|\n| 1 | 2 |\n"
    )
    out = _build_single(tmp_path, content)
    texts = [p.text for p in Document(out).paragraphs]
    assert "表 1-1 设备表" in texts, texts
    assert "表 1-2 材料表" in texts, texts


def test_docx_table_without_title_has_no_caption(tmp_path):
    """无表名行时不应插入空题注（避免出现光秃秃的「表 1-1」）"""
    content = "| A | B |\n|---|---|\n| 1 | 2 |\n"
    out = _build_single(tmp_path, content)
    texts = [p.text for p in Document(out).paragraphs]
    assert not any(t.startswith("表 1-") for t in texts), texts


# ===========================================================================
# AI 配图（文生图，与图表不同模态）：正文内嵌图片应作为真实位图进入 DOCX
# ===========================================================================
def _build_illustration_docx(tmp_path, content: str, image_bytes: dict) -> str:
    """构建含 AI 配图的单章节 DOCX（image_bytes 为空 dict 模拟下载失败）。"""
    sections = [{
        "id": "c1", "parent_id": "", "sort_order": 0, "level": 1,
        "title": "第一章 基坑支护", "content": content,
    }]
    blocks = {"c1": _parse_content_blocks(content)}
    out = str(tmp_path / "illu.docx")
    _build_docx_sync(
        out, {"id": "s1", "project_id": "p1", "name": "测试方案"},
        sections, {}, {}, {}, blocks,
        "宋体", 12, "", "", False, False, False,
        bidder_name="",
        heading_styles=_load_heading_styles({}),
        page_break_before_chapter=True,
        line_spacing=1.15,
        page_number_style="simple",
        toc_depth=3,
        margins=None, cover_info=None,
        image_bytes=image_bytes,
    )
    return out


def test_docx_content_subheadings_shift_child_section_numbers(tmp_path):
    """✅ E3（2026-09-25）：有 DB 子章节的章节，正文子标题降级为节内 body 命名空间。

    本节**内容**自带 "## 子标题" 且本节还有 DB 子章节时：
    - 旧实现（计数器右移续排）：正文 "1.1 项目概况/1.2 建筑概况"，DB 子章节 "1.3 工程规模"
    - ✅ 新实现（E3 降级隔离）：正文子标题走 body 命名空间 "1）项目概况/2）建筑概况"，
      与 DB 子章节的 X.X 命名空间彻底隔离——DB 子章节从 "1/2" 开始，不再被正文占用。
    两种手段都实现"正文子标题与 DB 子章节不撞号"这个语义目标。
    """
    content = ("本节概述。\n\n## 项目概况\n\n项目概况正文。\n\n"
               "## 建筑概况\n\n建筑概况正文。")
    sections = [
        {"id": "c1", "parent_id": "", "sort_order": 0, "level": 1,
         "title": "第一章 工程概况", "content": "第一章导语。"},
        {"id": "p1", "parent_id": "c1", "sort_order": 0, "level": 2,
         "title": "工程基本信息", "content": content},
        {"id": "s1", "parent_id": "p1", "sort_order": 0, "level": 3,
         "title": "工程规模与结构形式", "content": "规模正文。"},
        {"id": "s2", "parent_id": "p1", "sort_order": 1, "level": 3,
         "title": "脚手架搭设部位与高度", "content": "部位正文。"},
    ]
    blocks = {s["id"]: _parse_content_blocks(s["content"]) for s in sections}
    out = str(tmp_path / "shift.docx")
    _build_docx_sync(
        out, {"id": "s1", "project_id": "p1", "name": "测试方案"},
        [sections[0]],                      # roots：只含一级章，子节点走 children_map
        {"c1": [sections[1]], "p1": sections[2:]},
        {}, {}, blocks,
        "宋体", 12, "", "", False, False, False,
        bidder_name="",
        heading_styles=_load_heading_styles({}),
        page_break_before_chapter=True,
        line_spacing=1.15,
        page_number_style="simple",
    )
    joined = "\n".join(p.text for p in Document(out).paragraphs)
    # ✅ E3 降级后：正文子标题走 body 命名空间（1）、2）、…，不再占用 X.X 编号
    assert "1）、项目概况" in joined, joined
    assert "2）、建筑概况" in joined, joined
    # DB 子章节展示编号是累积路径（heading_v2：L2→1，L3→1.1 / 1.2）
    # —— 旧计数器前移手段让 DB 从 1.3/1.4 续排；E3 降级隔离后 DB 回到 1.1/1.2
    # （因为正文子标题已不再占用 1.1/1.2 编号），更符合章节层级直觉。
    assert "1.1 工程规模与结构形式" in joined, joined
    assert "1.2 脚手架搭设部位与高度" in joined, joined
    # ✅ 核心验证：正文子标题不再占用 1.1/1.2 编号 → 成稿无撞号
    assert "1.1 项目概况" not in joined, joined
    assert "1.2 建筑概况" not in joined, joined


def test_docx_embeds_ai_illustration(tmp_path):
    """正文里的 ![说明](http...) 应作为真实位图嵌入，并生成「图 X-Y 说明」图题"""
    url = "http://example.com/illu.png"
    content = f"基坑支护施工说明。\n\n![基坑支护示意]({url})\n"
    out = _build_illustration_docx(tmp_path, content, {url: BytesIO(_png_bytes())})
    with zipfile.ZipFile(out) as z:
        assert [n for n in z.namelist() if n.startswith("word/media/")], "配图未嵌入 DOCX"
    paras = [p.text for p in Document(out).paragraphs]
    assert "图 1-1 基坑支护示意" in paras, paras
    # 裸 Markdown 图片语法不应残留在成稿里
    assert not any("![" in t for t in paras), paras


def test_docx_illustration_without_bytes_keeps_caption(tmp_path):
    """配图下载失败（image_bytes 未命中）时：整块跳过、**不占图号**、不留孤立图题。

    ✅ 2026-09-26（陈旧断言更新）：旧实现 `_add_illustration_from_bytes` 只写
    图题不写位图，但 figure_counters 已 +1 → 交付文档出现「图 1-1」缺失、
    编号却从「图 1-2」起跳的**图号虚跳**。现与图表分支同口径修复：
    先判可用性（`_chart_ok`），不可用则不占图号、不留孤立图题，
    并回收孤儿引导语（见 export.py:3403-3412 的注释）。
    本用例据此改为断言「无图题、无图号、无报错文本残留」。

    注意：函数名沿用历史命名（keeps_caption），实际语义已变为
    "skips_and_frees_fig_number"。
    """
    content = "基坑支护施工说明。\n\n![基坑支护示意](http://example.com/missing.png)\n"
    out = _build_illustration_docx(tmp_path, content, {})
    paras = [p.text for p in Document(out).paragraphs]
    # 不应残留图题（否则占用图号造成后续编号虚跳）
    assert not any("基坑支护示意" in t for t in paras), paras
    assert not any("图 1-1" in t for t in paras), paras
    # 也不得出现渲染失败类报错文本
    assert not any("渲染失败" in t or "插入失败" in t for t in paras), paras
    # 正文本体仍应正常写入
    assert any("基坑支护施工说明" in t for t in paras), paras


def test_docx_figure_number_not_skipped_by_failed_illustration(tmp_path):
    """配图失败不占图号 → 后续成功配图从「图 1-1」起（不再虚跳到 1-2）。"""
    bad = "http://example.com/missing.png"
    good = "http://example.com/ok.png"
    content = (
        "基坑支护施工说明。\n\n"
        f"![基坑支护示意]({bad})\n\n"
        f"![监测布置示意]({good})\n"
    )
    out = _build_illustration_docx(
        tmp_path, content, {good: BytesIO(_png_bytes())})
    paras = [p.text for p in Document(out).paragraphs]
    assert "图 1-1 监测布置示意" in paras, paras
    assert not any("基坑支护示意" in t for t in paras), paras


def test_docx_ai_image_placeholder_skipped_without_leaking_prompt(tmp_path):
    """未生成的 ```ai_image 占位块：导出整块跳过，绝不把内部 JSON 印进成稿。

    BUG 回归（2026-09-17）：旧实现把 ```ai_image 当普通代码块走 `_add_code_block` ——
    `{"prompt": "...", "style": "engineering_diagram", "title": "..."}` 原样出现在
    交付 DOCX 里（AI 绘图提示词泄漏，且观感极差）。
    现改为：跳过该块 + 不写红字占位 + 不占用图号（用户侧由导出预检的
    chart_ungenerated 告警提示"请先生成配图"）。
    """
    content = (
        "本节配图示意如下：\n\n"
        "```ai_image\n"
        '{"prompt": "深基坑支护结构剖面，含支护桩、锚索、冠梁", '
        '"style": "engineering_diagram", "title": "支护剖面图"}\n'
        "```\n\n"
        "后续正文段落。\n"
    )
    out = _build_illustration_docx(tmp_path, content, {})
    paras = [p.text for p in Document(out).paragraphs]
    joined = "\n".join(paras)

    # 正文照常保留
    assert "本节配图示意如下：" in joined and "后续正文段落。" in joined, paras
    # 1) 内部 JSON / 绘图提示词不得泄漏
    assert "prompt" not in joined, paras
    assert "engineering_diagram" not in joined, paras
    assert "支护桩" not in joined, paras
    # 2) 不出现红字占位文本
    assert "渲染失败" not in joined and "插入失败" not in joined, paras
    # 3) 不占用图号（无图即无图题）
    assert not any(t.startswith("图 ") for t in paras), paras
    # 4) 不得嵌入任何位图
    with zipfile.ZipFile(out) as z:
        assert not [n for n in z.namelist() if n.startswith("word/media/")], paras


