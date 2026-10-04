# -*- coding: utf-8 -*-
"""图表生成与渲染模块深度修复回归（2026-09-25）。

覆盖三类缺陷（每类都带"修复前行为"的显式反例断言）：

① 图片尺寸溢出页面（P0，export.py）
   渲染器产出 2.5x 超采样栅格图（后端 PIL + 300 DPI 元数据 / 前端 canvas），
   旧实现只按 `px × 2.54 / 96` 估宽、且**只用 16cm 宽度封顶、完全无视高度**：
   纵向流程图（`flowchart TD` 十几个节点，专项方案里极常见）实测 700×4070px →
   宽度封顶 16cm、高度 16×4070/700 = **93.03cm**（页面可用正文高 Letter 22.94cm /
   A4 24.7cm）→ 成稿单张图溢出页面、被裁切。
   修复：`_fit_image_cm` 用 PNG DPI 换算真实物理尺寸 + 「16cm × 22cm」双上限等比缩放。

② 长载荷渲染缓存串图（P1，mermaid_renderer.py）
   旧 `_render_cache_key` 用 `str(code)[:8192]` 做键 → 前 8192 字符相同的两张
   不同图表共用缓存槽，后一张直接命中前一张的 PNG（张冠李戴且完全静默）。

③ 删图留孤儿引导语（P1，_chart_pipeline.py + export.py）
   图表块因校验修复失败 / 超配图上限被删，或导出时渲染失败被跳过时，
   专为引出该图而写的"施工工艺流程如下图所示："仍留在正文 →「见下图」却无图。
   修复：删块/跳过时同步回收该引导语（管线侧与导出侧同口径）。
"""
from __future__ import annotations

import asyncio
import io
import json
import os
import re
import sys
import zipfile

from docx import Document
from PIL import Image

sys.path.insert(0, r"J:\编程\专项方案工具箱\backend")

from app.routers._chart_pipeline import (  # noqa: E402
    _drop_dangling_lead_in,
    register_inline_charts,
)
from app.routers.export import (  # noqa: E402
    _CHART_MAX_HEIGHT_CM,
    _CHART_MAX_WIDTH_CM,
    _DEFAULT_HEADING_STYLES,
    _build_docx_sync,
    _fit_image_cm,
    _read_image_dpi,
)
from app.services.ai.mermaid_renderer import (  # noqa: E402
    _code_cache_ident,
    _render_cache_key,
)

# 1 cm = 360000 EMU（docx 内部长度单位）
_EMU_PER_CM = 360000


# ---------------------------------------------------------------------------
# 工具
# ---------------------------------------------------------------------------

def _png(w: int, h: int, dpi: tuple | None = None) -> bytes:
    """生成指定像素尺寸（可选 DPI 元数据）的 PNG 字节流。"""
    buf = io.BytesIO()
    img = Image.new("RGB", (w, h), (210, 225, 245))
    if dpi:
        img.save(buf, format="PNG", dpi=dpi)
    else:
        img.save(buf, format="PNG")
    return buf.getvalue()


def _docx_drawing_extents(path: str) -> list[tuple[float, float]]:
    """读取 DOCX 中每张内嵌图片的**显示尺寸**（cm）：``[(宽, 高), ...]``。"""
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    return [(int(m.group(1)) / _EMU_PER_CM, int(m.group(2)) / _EMU_PER_CM)
            for m in re.finditer(r'<wp:extent\s+cx="(\d+)"\s+cy="(\d+)"', xml)]


class FakeDb:
    """最小异步 db mock（与 test_inline_charts 同款，仅记录调用）。

    ✅ F-3 修复配套（2026-10-05）：真实 aiosqlite 上
      · SELECT 成功返回 Cursor（有 .fetchall()）；异常时才返回 None
      · INSERT 成功返回 Cursor（有 .rowcount）；异常时才返回 None
    测试桩必须与真实语义一致，否则 F-3 修复（INSERT 返回 None →
    降级为 skipped → 正文被裁剪）会让历史断言误判为「图块被删」。
    """

    def __init__(self):
        self.rows = []
        self.commits = 0

    async def execute(self, sql, params=None):
        self.rows.append((sql, tuple(params or ())))
        head = sql.lstrip().upper()
        if head.startswith("SELECT"):
            class _C:
                async def fetchall(self):
                    return []
            return _C()
        # INSERT / DELETE 均返回 Cursor 语义
        class _FakeCur:
            rowcount = 1
            lastrowid = 1
        return _FakeCur()

    async def commit(self):
        self.commits += 1


def _sec(sid, order, title, content):
    return {"id": sid, "parent_id": "", "level": 1,
            "sort_order": order, "title": title, "content": content}


def _build(out, sections, rendered_bytes, chart_lookup=None):
    """用最小 prep 组装 DOCX（与 test_export_fallback_dedup 同款调用口径）。"""
    _build_docx_sync(**{
        "out_path": out, "scheme": {"name": "回归", "project_id": "P1"},
        "roots": sections, "children_map": {"": sections},
        "chart_lookup": chart_lookup or {},
        "rendered_bytes": rendered_bytes,
        "blocks_cache": {}, "font_name": "宋体", "font_size": 12,
        "page_header": "", "page_footer": "", "show_page_number": False,
        "show_title_page": False, "show_toc": False, "bidder_name": "",
        "heading_styles": _DEFAULT_HEADING_STYLES,
        "page_break_before_chapter": True, "line_spacing": 1.15,
        "page_number_style": "simple", "toc_depth": 3,
        "margins": None, "cover_info": None, "image_bytes": {},
        "global_facts": [], "chart_fail_placeholder": False,
    })
    return Document(out)


def _doc_text(path):
    """DOCX 全部段落文本（换行拼接）。"""
    return "\n".join(p.text for p in Document(path).paragraphs)


def _run(coro):
    return asyncio.run(coro)


# ===========================================================================
# ① 图片尺寸：DPI 换算 + 宽/高双上限
# ===========================================================================

def test_fit_image_cm_wide_chart_keeps_column_width():
    """宽图（后端实测 flowchart 3300×710@300dpi）仍按 16cm 栏宽插入（无回归）。"""
    w_cm, h_cm = _fit_image_cm(3300, 710, 300)
    assert w_cm == _CHART_MAX_WIDTH_CM == 16.0
    # 与旧算法（px×2.54/96 后被 16cm 封顶）产出的宽度完全一致
    assert min(3300 * 2.54 / 96.0, 16.0) == w_cm
    assert h_cm == round(16.0 * 710 / 3300, 2) == 3.44


def test_fit_image_cm_tall_chart_clamped_by_page_height():
    """纵向长图（700×4070@300dpi）：高度必须落在页面内；旧算法产出 93cm。"""
    w_cm, h_cm = _fit_image_cm(700, 4070, 300)
    legacy_w = min(700 * 2.54 / 96.0, 16.0)
    legacy_h = legacy_w * 4070 / 700
    assert legacy_h > 90.0, "反例：旧算法确实产出 93cm 高的图（页面仅 22.94cm）"
    assert h_cm == _CHART_MAX_HEIGHT_CM == 22.0
    assert abs(w_cm - round(700 * 2.54 / 300 * (22.0 / (4070 * 2.54 / 300)), 2)) < 0.02
    assert w_cm < 4.0, "高度封顶后宽度随之等比缩小"


def test_fit_image_cm_dpi_fallback_and_small_image_not_upscaled():
    """无 DPI 元数据（前端 canvas PNG）回退 96dpi；小图不放大。"""
    assert _fit_image_cm(600, 400, None) == _fit_image_cm(600, 400, 0) \
        == (round(600 * 2.54 / 96.0, 2), round(400 * 2.54 / 96.0, 2))
    # 400×200@300dpi 物理尺寸很小 → 保持原尺寸，不拉伸到栏宽
    w_cm, h_cm = _fit_image_cm(400, 200, 300)
    assert w_cm == round(400 * 2.54 / 300, 2) and h_cm == round(200 * 2.54 / 300, 2)
    assert w_cm < _CHART_MAX_WIDTH_CM


def test_fit_image_cm_invalid_dims_fallback_width_only():
    """像素尺寸缺失/非法 → 退回"按栏宽插入、不声明高度"。"""
    assert _fit_image_cm(0, 0, 300) == (_CHART_MAX_WIDTH_CM, 0.0)
    assert _fit_image_cm("x", "y", 300) == (_CHART_MAX_WIDTH_CM, 0.0)
    assert _fit_image_cm(None, None, None) == (_CHART_MAX_WIDTH_CM, 0.0)


def test_read_image_dpi_tolerates_dirty_metadata():
    """DPI 元数据容错：元组 / 标量 / 字符串 / 脏值 / 缺失。"""
    assert _read_image_dpi((300.0, 300.0)) == 300.0
    assert _read_image_dpi(300) == 300.0
    assert _read_image_dpi("300") == 300.0
    assert _read_image_dpi(None) == 0.0
    assert _read_image_dpi((0, 0)) == 0.0
    assert _read_image_dpi("bad") == 0.0
    assert _read_image_dpi([]) == 0.0


def test_docx_chart_sizes_end_to_end(tmp_path):
    """端到端：宽图 16cm 栏宽；纵向长图不溢出页面；图题规范「图 X-Y 图名」。"""
    wide = "flowchart LR\n  A[\"施工准备\"] --> B[\"成品保护\"]"
    tall = "flowchart TD\n  A[\"施工准备\"] --> B[\"基层清理\"]\n  B --> C[\"验收\"]"
    sections = [
        _sec("s1", 0, "第一章 工艺流程",
             "工艺流程如下图所示：\n\n```mermaid\n" + wide + "\n```\n"),
        _sec("s2", 1, "第二章 竖向流程",
             "竖向流程如下图所示：\n\n```mermaid\n" + tall + "\n```\n"),
    ]
    rendered = {
        ("flowchart", wide): io.BytesIO(_png(3300, 710, (300, 300))),
        ("flowchart", tall): io.BytesIO(_png(700, 4070, (300, 300))),
    }
    out = os.path.join(str(tmp_path), "size.docx")
    doc = _build(out, sections, rendered)
    extents = _docx_drawing_extents(out)
    assert len(extents) == 2, extents
    wide_w, wide_h = extents[0]
    tall_w, tall_h = extents[1]
    assert abs(wide_w - 16.0) < 0.05, f"宽图应保持栏宽 16cm，实际 {wide_w}"
    assert abs(wide_h - 3.44) < 0.05, f"宽图高度应按比例，实际 {wide_h}"
    assert tall_h <= _CHART_MAX_HEIGHT_CM + 0.05, f"长图不得溢出页面，实际 {tall_h}"
    assert tall_w < 5.0, f"长图高度封顶后宽度等比缩小，实际 {tall_w}"
    captions = [p.text.strip() for p in doc.paragraphs
                if p.text.strip().startswith("图 ")]
    assert captions == ["图 1-1 工艺流程", "图 2-1 竖向流程"], captions


# ===========================================================================
# ② 长载荷渲染缓存键（不得截断）
# ===========================================================================

def test_render_cache_key_long_payload_not_truncated():
    """反例：9000 字符前缀相同的两张不同图表，旧实现键完全相同（串图）。"""
    prefix = '{"type":"labor","title":"劳动力配置计划","phases":[' + "A" * 8900
    code_a = prefix + '],"data":[[1,2]]}'
    code_b = prefix + '],"data":[[3,4]]}'
    assert code_a[:8192] == code_b[:8192], "反例前提：前 8192 字符相同"
    assert code_a != code_b
    key_a = _render_cache_key(code_a, "labor", 90, "天", 5, False, True)
    key_b = _render_cache_key(code_b, "labor", 90, "天", 5, False, True)
    assert key_a != key_b, "长载荷不得因截断共用缓存键（否则命中上一张图）"


def test_render_cache_key_short_payload_keeps_legacy_form():
    """短载荷（≤8192）键值保持原样 —— 既有磁盘缓存与单测口径不变。"""
    code = "graph TD; A-->B;"
    assert _code_cache_ident(code) == code
    # ✅ 2026-09-27：键首段新增渲染器版本号（memory LRU 与磁盘缓存同步失效）
    from app.services.ai import mermaid_renderer as _mr
    assert _render_cache_key(code, "flowchart", 90, "天", 5, False, True) == \
        (f"v{_mr._RENDERER_VERSION}", "flowchart", code, 90, "天", 5, False, True)
    assert _code_cache_ident(None) == ""


# ===========================================================================
# ③ 孤儿引导语 —— 管线侧（删块同时删引导语）
# ===========================================================================

def test_pipeline_removes_orphan_lead_in_on_bad_block():
    """坏块被删 → 其专属引导语一并移除，正文其余部分保留。"""
    db = FakeDb()
    content = ("本章说明施工顺序。\n"
               "施工工艺流程如下图所示：\n\n"
               "```mermaid\n"
               "graph TD\n"
               "A-[坏\n"
               "```\n"
               "后文继续说明。")
    n, new_content = _run(register_inline_charts(db, "scheme-1", "sec-l1", content))
    assert n == 0
    assert "```mermaid" not in new_content
    assert "如下图所示" not in new_content, "孤儿引导语必须一并移除"
    assert "本章说明施工顺序。" in new_content
    assert "后文继续说明。" in new_content


def test_pipeline_keeps_plain_prose_before_bad_block():
    """回归锁定：删块不得误删非引导语正文（旧行为不变）。"""
    db = FakeDb()
    content = "前文说明\n```mermaid\ngraph TD\nA-[坏\n```\n后文说明"
    n, new_content = _run(register_inline_charts(db, "scheme-1", "sec-l2", content))
    assert n == 0
    assert "前文说明" in new_content and "后文说明" in new_content
    assert "```mermaid" not in new_content


def test_pipeline_drop_dangling_lead_in_scope_guard():
    """引导语识别口径：只认"图/表"类引导尾；不碰标题/列表/长句。"""
    out = ["## 施工工艺", ""]
    assert _drop_dangling_lead_in(out) is False
    out = ["各阶段投入如下：", ""]
    assert _drop_dangling_lead_in(out) is False, "无图/表关键词的引导句不得误删"
    out = ["- 施工流程如下图所示：", ""]
    assert _drop_dangling_lead_in(out) is False, "无序列表项不得被当作引导语段落"
    out = ["1. 施工流程如下图所示：", ""]
    assert _drop_dangling_lead_in(out) is False, "有序列表项不得被当作引导语段落"
    out = ["（1）施工流程如下图所示：", ""]
    assert _drop_dangling_lead_in(out) is False, "中文编号列表项不得被当作引导语段落"
    out = ["5.2 施工流程如下图所示：", ""]
    assert _drop_dangling_lead_in(out) is False, "点分编号小标题不得被删"
    out = ["| 施工流程如下图所示： |", ""]
    assert _drop_dangling_lead_in(out) is False, "表格行不得被删"
    out = ["x" * 70 + "施工流程如下图所示：", ""]
    assert _drop_dangling_lead_in(out) is False, "长句不得被删"
    out = ["正文段落。", "施工工艺流程如下图所示：", "", ""]
    assert _drop_dangling_lead_in(out) is True
    assert out == ["正文段落。"], out
    # 引导语与围栏之间没有空行时同样能识别
    out = ["正文段落。", "施工工艺流程如下图。"]
    assert _drop_dangling_lead_in(out) is True
    assert out == ["正文段落。"], out


def test_pipeline_limit_trim_also_drops_second_lead_in():
    """每章≤1 上限裁剪第 2 块时，其引导语同样回收（第 1 块的保留）。"""
    db = FakeDb()
    first = "flowchart TD\n    A --> B\n    B --> C"
    second = "flowchart LR\n    C --> D"
    content = ("第一张流程如下图所示：\n\n"
               "```mermaid\n" + first + "\n```\n\n"
               "第二张流程如下图所示：\n\n"
               "```mermaid\n" + second + "\n```\n")
    n, new_content = _run(register_inline_charts(
        db, "scheme-1", "sec-l3", content, enforce_limits=True))
    assert n == 1
    assert first in new_content
    assert second not in new_content
    assert new_content.count("如下图所示") == 1, new_content
    assert "第一张流程如下图所示：" in new_content



# ===========================================================================
# ③ 孤儿引导语 —— 导出侧（跳过图 → 回收段落；正常出图不得误删）
# ===========================================================================

def test_docx_drops_orphan_lead_in_when_render_failed(tmp_path):
    """渲染失败被跳过 → 图与引导语都不进成稿，且不占图号。"""
    code = "flowchart TD\n  A[\"施工准备\"] --> B[\"基层清理\"]\n"
    content = ("本章说明施工顺序。\n\n"
               "施工工艺流程如下图所示：\n\n"
               "```mermaid\n" + code + "\n```\n\n"
               "后续说明文字。\n")
    out = os.path.join(str(tmp_path), "orphan.docx")
    _build(out, [_sec("s1", 0, "第一章 施工工艺", content)],
           {})                                   # 渲染失败（无字节）
    text = _doc_text(out)
    assert "如下图所示" not in text, "无图就不该留「见下图」的悬空引用"
    assert "图 1-1" not in text
    assert "本章说明施工顺序。" in text and "后续说明文字。" in text


def test_docx_keeps_lead_in_when_chart_rendered(tmp_path):
    """正常出图时引导语与图题都必须保留（不得误删）。"""
    code = "flowchart LR\n  A[\"施工准备\"] --> B[\"成品保护\"]"
    content = "施工工艺流程如下图所示：\n\n```mermaid\n" + code + "\n```\n"
    out = os.path.join(str(tmp_path), "keep.docx")
    doc = _build(out, [_sec("s1", 0, "第一章 施工工艺", content)],
                 {("flowchart", code): io.BytesIO(_png(1200, 400, (300, 300)))})
    text = _doc_text(out)
    assert "施工工艺流程如下图所示：" in text
    assert len(doc.inline_shapes) == 1
    captions = [p.text.strip() for p in doc.paragraphs
                if p.text.strip().startswith("图 ")]
    assert captions == ["图 1-1 施工工艺流程"], captions


def test_docx_does_not_drop_prose_when_chart_skipped(tmp_path):
    """范围守卫：跳过图时不得误删普通正文段落。"""
    code = "flowchart TD\n  A[\"施工准备\"] --> B[\"基层清理\"]\n"
    content = "各阶段投入如下：\n\n```mermaid\n" + code + "\n```\n"
    out = os.path.join(str(tmp_path), "guard.docx")
    _build(out, [_sec("s1", 0, "第一章 施工工艺", content)], {})
    assert "各阶段投入如下：" in _doc_text(out), "非图/表引导句的正文不得被删"


def test_chart_payload_envelope_roundtrip_untouched():
    """回归：本次修复不改变图表载荷信封契约（唯一构造器/解析器互逆）。"""
    from app.services.chart_payload import (
        build_chart_envelope,
        extract_chart_payload,
        is_canonical_chart_payload,
    )
    env = build_chart_envelope(code="graph TD; A-->B", title="施工流程")
    assert json.loads(env)["mermaid_code"] == "graph TD; A-->B"
    assert extract_chart_payload(env) == "graph TD; A-->B"
    assert is_canonical_chart_payload(env)
