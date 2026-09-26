"""回归测试：兜底图表跨章节去重（2026-09-20 深度审查修复）。

缺陷：旧实现兜底渲染去重键为 ("fallback", chart_type, "")，与具体章节无关 ——
多个章节的同类型图表都无注册代码时，只有第一张被渲染，其余被 rendered_charts
静默丢弃（图号缺失、无系统侧信号）。修复后键含 (sec_id, "fallback", chart_type,
_norm_code(code))，每章各自渲染。

场景：三章均含 [CHART_TYPE: labor] 标记；仅 secC 注册了代码；secA/secB 走兜底。
期望：三张图（图 1-1 / 2-1 / 3-1）全部渲染。
"""
import io
import os
import sys
import tempfile

import pytest
from PIL import Image
from docx import Document

sys.path.insert(0, r"J:\编程\专项方案工具箱\backend")

from app.routers.export import _build_docx_sync, _DEFAULT_HEADING_STYLES

_buf = io.BytesIO()
Image.new("RGB", (200, 100), (200, 30, 30)).save(_buf, "PNG")
PNG_VALID = _buf.getvalue()

CODE_LABOR = "graph TD\n  A[劳动力计划] --> B[进场安排]\n"


def _make_sec(sid, order, title, content):
    return {"id": sid, "parent_id": "", "level": 1,
            "sort_order": order, "title": title, "content": content}


def _build(secs, chart_lookup):
    out = os.path.join(tempfile.mkdtemp(prefix="fb_dedup_"), "out.docx")
    _build_docx_sync(**{
        "out_path": out, "scheme": {"name": "回归", "project_id": "P1"},
        "roots": secs, "children_map": {"": secs},
        "chart_lookup": chart_lookup,
        "rendered_bytes": {("labor", CODE_LABOR): io.BytesIO(PNG_VALID)},
        "blocks_cache": {}, "font_name": "宋体", "font_size": 12,
        "page_header": "", "page_footer": "", "show_page_number": False,
        "show_title_page": False, "show_toc": False, "bidder_name": "",
        "heading_styles": _DEFAULT_HEADING_STYLES,
        "page_break_before_chapter": True, "line_spacing": 1.15,
        "page_number_style": "simple", "toc_depth": 3,
        "margins": None, "cover_info": None, "image_bytes": {},
        "global_facts": [], "chart_fail_placeholder": False,
    })
    return Document(out)


def test_fallback_chart_multi_section_not_dropped():
    secs = [
        _make_sec("secC", 0, "第一章 劳动力组织", "## 总述\n[CHART_TYPE: labor]"),
        _make_sec("secA", 1, "第二章 劳动力保障", "## 概述\n[CHART_TYPE: labor]"),
        _make_sec("secB", 2, "第三章 劳动力培训", "## 安排\n[CHART_TYPE: labor]"),
    ]
    doc = _build(secs, {("secC", "labor"): CODE_LABOR})
    captions = [p.text.strip() for p in doc.paragraphs
                if p.text.strip().startswith("图 ")]
    # 三个章节各一张 → 3 张图 + 3 条图注
    assert len(doc.inline_shapes) == 3, f"实际 {len(doc.inline_shapes)} 张"
    assert captions == ["图 1-1 劳动力配置计划",
                        "图 2-1 劳动力配置计划",
                        "图 3-1 劳动力配置计划"], captions


def test_inline_chart_same_type_multiple_in_one_section():
    """同章节内两张不同代码的同类内联图表不应互相吞并（既有语义保持）。"""
    code2 = "graph TD\n  C --> D\n"
    secs = [_make_sec("s1", 0, "第一章 A", "## 总述\n"
                      "```mermaid\n" + CODE_LABOR + "\n```\n"
                      "```mermaid\n" + code2 + "\n```")]
    from app.routers.export import _parse_content_blocks
    blocks = _parse_content_blocks(secs[0]["content"])
    chart_codes = [b["code"] for b in blocks if b["type"] == "chart"]
    assert len(chart_codes) == 2, chart_codes
    out = os.path.join(tempfile.mkdtemp(prefix="fb_inline_"), "out.docx")
    _build_docx_sync(**{
        "out_path": out, "scheme": {"name": "回归", "project_id": "P1"},
        "roots": secs, "children_map": {"": secs},
        "chart_lookup": {},
        # 键必须与 write_section 实际使用的解析后代码一致
        "rendered_bytes": {("flowchart", chart_codes[0]): io.BytesIO(PNG_VALID),
                           ("flowchart", chart_codes[1]): io.BytesIO(PNG_VALID)},
        "blocks_cache": {}, "font_name": "宋体", "font_size": 12,
        "page_header": "", "page_footer": "", "show_page_number": False,
        "show_title_page": False, "show_toc": False, "bidder_name": "",
        "heading_styles": _DEFAULT_HEADING_STYLES,
        "page_break_before_chapter": True, "line_spacing": 1.15,
        "page_number_style": "simple", "toc_depth": 3,
        "margins": None, "cover_info": None, "image_bytes": {},
        "global_facts": [], "chart_fail_placeholder": False,
    })
    doc = Document(out)
    assert len(doc.inline_shapes) == 2, f"实际 {len(doc.inline_shapes)} 张"
