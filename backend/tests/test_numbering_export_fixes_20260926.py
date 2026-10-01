"""编号规则跨模块统一（2026-09-26）回归测试。

覆盖用户给定导出标准与三模块编号一致性的关键约束：
  - D2：正文子标题（节内 body 命名空间）L6/L7 格式必须与导出行级标题同口径
        （1）、标题 / a、标题，顿号、无空格）；4 段编号（L5 等价）带顿号。
  - D3：导出 DOCX 所有标题（L1~L7）不得倾斜，含标题内嵌 Markdown 行内斜体兜底。
  - D4：章节增删/顺序调整（reorder/create/delete/update 结构变更）后，已落库正文的
        子标题编号按新编号统一重规范化（落库正文 = 导出成稿 同源）。
"""
import json

import pytest


# ============================================================
# D2：正文子标题编号格式（与导出行级标题同口径）
# ============================================================
class TestBodySubheadingFormatAlignment:
    async def test_demote_l6_uses_paren_punctuation(self):
        """有 DB 子章节时，一级正文子标题应为 '1）、标题'（顿号、无空格）。"""
        from app.routers.export import _compute_subheading
        sub = {}
        text, style = _compute_subheading("2", 2, 1, sub, "总体安排",
                                         sec_id="s", has_children=True)
        assert text == "1）、总体安排", text
        assert style == 6
        # 计数器续排
        text2, _ = _compute_subheading("2", 2, 1, sub, "资源配置",
                                      sec_id="s", has_children=True)
        assert text2 == "2）、资源配置", text2

    async def test_demote_l7_uses_punctuation_no_space(self):
        """更深层正文子标题应为 'a、标题'（顿号、无空格）。"""
        from app.routers.export import _compute_subheading
        sub = {}
        _compute_subheading("2", 2, 1, sub, "总体安排", has_children=True)  # rel0 → 1）
        text, style = _compute_subheading("2", 2, 2, sub, "劳动组织",
                                          sec_id="s", has_children=True)  # rel1
        assert text == "a、劳动组织", text
        assert style == 7
        # 字母序列续排
        _compute_subheading("2", 2, 3, sub, "工种", sec_id="s", has_children=True)
        text3, _ = _compute_subheading("2", 2, 4, sub, "调度",
                                      sec_id="s", has_children=True)
        assert text3 == "c、调度", text3

    async def test_nodemote_three_seg_space_and_four_seg_punctuation(self):
        """无 DB 子章节：3 段编号用空格（L4 等价），4 段编号用顿号（L5 等价）。

        base 由首次调用决定（md_lv=2 → base=2），需按 2→3→4 顺序累加相对深度才能
        推到 4 段（rel=2 → 1.1.1.1）。
        """
        from app.routers.export import _compute_subheading
        sub = {}
        t3, _ = _compute_subheading("1", 2, 2, sub, "总体安排",
                                   sec_id="s", has_children=False)  # rel0 → 1.1
        assert t3 == "1.1 总体安排", t3
        t3b, _ = _compute_subheading("1", 2, 3, sub, "关键节点",
                                     sec_id="s", has_children=False)  # rel1 → 1.1.1
        assert t3b == "1.1.1 关键节点", t3b
        # 第四层：1.1.1.1（4 段）→ 顿号（对齐导出 L5 '1.1.1.1、'）
        t4, style4 = _compute_subheading("1", 2, 4, sub, "更深层",
                                         sec_id="s", has_children=False)  # rel2 → 1.1.1.1
        assert t4 == "1.1.1.1、更深层", t4
        assert style4 == 5

    async def test_format_matches_export_heading_standard(self):
        """正文子标题 L6/L7 格式必须与 heading_v2 / HEADING_STYLE_CONFIG 同口径。"""
        from app.routers.export import _compute_subheading
        from app.services.ai.heading_templates import HEADING_STYLE_CONFIG
        sub = {}
        text, style = _compute_subheading("2", 2, 1, sub, "总体安排",
                                         sec_id="s", has_children=True)
        assert text.endswith(HEADING_STYLE_CONFIG[style]["punctuation"] + "总体安排")


# ============================================================
# D3：导出标题不得倾斜（含内联 Markdown 斜体兜底）
# ============================================================
class TestHeadingNonItalic:
    def test_finalize_kills_inline_italic(self):
        """标题内嵌 *斜体* 经行内渲染后局部倾斜，_finalize_heading_runs 必须全部复位。"""
        from docx import Document
        from app.routers.export import (_add_runs_with_inline_format,
                                        _finalize_heading_runs)
        p = Document().add_paragraph()
        _add_runs_with_inline_format(p, "第一章 *编制* 说明")
        # 前置：内联斜体确已被行内渲染设为 italic
        assert any(r.italic for r in p.runs), "前置条件：内联 *斜体* 应产生倾斜 run"
        _finalize_heading_runs(p)
        assert all(not r.italic for r in p.runs), \
            "标题所有 run 必须非倾斜：" + str([r.italic for r in p.runs])

    def test_finalize_on_explicit_italic_runs(self):
        """直接构造倾斜 run 也应被强制复位（验证兜底逻辑本身）。"""
        from docx import Document
        from app.routers.export import _finalize_heading_runs
        p = Document().add_paragraph()
        r = p.add_run("第一章 编制综合说明")
        r.italic = True
        _finalize_heading_runs(p)
        assert all(not x.italic for x in p.runs)


# ============================================================
# D4：结构变更后正文子标题按新编号重规范化
# ============================================================
async def _seed_scheme(db, pid="p1", sid="s1"):
    await db.execute("INSERT OR IGNORE INTO projects (id, name) VALUES (?,?)", (pid, "P"))
    await db.execute(
        "INSERT INTO schemes (id, project_id, name) VALUES (?,?,?)", (sid, pid, "S"))
    await db.commit()


class TestStructuralChangeRenormalize:
    async def test_reorder_renormalizes_sibling_content(self, db_conn):
        """拖拽重排后，所有含正文章节的子标题按新编号统一重规范化（D4）。"""
        from app.routers.sections import reorder_sections
        await _seed_scheme(db_conn)
        # sec3 原有 AI 写死的旧号 3.1；sec1 旧号 1.1
        await db_conn.execute(
            "INSERT INTO sections (id, scheme_id, project_id, title, parent_id, "
            "level, sort_order, outline_json, content) VALUES "
            "('sec1','s1','p1','工程概况','','1',0,?, '## 1.1 概况\n正文\n')",
            (json.dumps({"id": "1", "level": 1}, ensure_ascii=False),))
        await db_conn.execute(
            "INSERT INTO sections (id, scheme_id, project_id, title, parent_id, "
            "level, sort_order, outline_json, content) VALUES "
            "('sec3','s1','p1','施工部署','','1',1,?, '## 3.1 总体安排\n正文三\n')",
            (json.dumps({"id": "3", "level": 1}, ensure_ascii=False),))
        await db_conn.commit()

        # 重排：sec3 → 第 1 章，sec1 → 第 2 章
        await reorder_sections("s1", {"order": ["sec3", "sec1"]}, db=db_conn)

        cur = await db_conn.execute("SELECT content FROM sections WHERE id='sec3'")
        c3 = (await cur.fetchone())["content"]
        cur = await db_conn.execute("SELECT content FROM sections WHERE id='sec1'")
        c1 = (await cur.fetchone())["content"]
        # 重排后 sec3 为第 1 章 → 正文子标题应为 1.1 而非旧号 3.1
        assert "## 1.1 总体安排" in c3, c3
        assert "3.1" not in c3, c3
        # sec1 顺移为第 2 章 → 1.1 应重算为 2.1
        assert "## 2.1 概况" in c1, c1
        assert "1.1 概况" not in c1, c1

    async def test_renormalize_helper_idempotent_and_safe(self, db_conn):
        """_renormalize_all_section_contents 幂等且单章失败不阻断事务。"""
        from app.routers.sections import _renormalize_all_section_contents
        await _seed_scheme(db_conn)
        await db_conn.execute(
            "INSERT INTO sections (id, scheme_id, project_id, title, parent_id, "
            "level, sort_order, outline_json, content) VALUES "
            "('sec1','s1','p1','工程概况','','1',0,?, '## 1.1 概况\n正文\n')",
            (json.dumps({"id": "1", "level": 1}, ensure_ascii=False),))
        await db_conn.commit()
        await _renormalize_all_section_contents(db_conn, "s1")
        cur = await db_conn.execute("SELECT content FROM sections WHERE id='sec1'")
        first = (await cur.fetchone())["content"]
        # 再跑一次应无变化（幂等）
        await _renormalize_all_section_contents(db_conn, "s1")
        cur = await db_conn.execute("SELECT content FROM sections WHERE id='sec1'")
        second = (await cur.fetchone())["content"]
        assert first == second
