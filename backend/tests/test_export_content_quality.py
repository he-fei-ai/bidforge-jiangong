# -*- coding: utf-8 -*-
"""导出内容质量回归（2026-09-19）。

背景（全部基于**真实交付文档取证**，见 _diagnostics/_diag_docx_output.py）：
1. 正文写「脚手架分段搭设与转换顺序如下图所示：」，紧跟的图题却是
   「图 4-1 劳动力配置计划」—— 导出图题一律取「图表类型通用名」，与本章语义错配；
2. chart-json 载荷缺 ``type`` 键时被静默兜底为 ``labor``，labor 渲染器对该结构
   必然返回 None → 交付文档出现红色「[图 X-Y … — 渲染失败]」占位；
3. 生成被截断（``flowchart LR`` 只写到 ``B --> C{``）时围栏未闭合，残片既渲染
   失败、又可能以源码形式印进成稿；
4. 正文里 73 处 ``××`` 占位符、19 处「招标」类投标用语无任何机器可枚举的清单。

本模块把这些口径锁成不变量。
"""
import json

import pytest

# ---------------------------------------------------------------------------
# ① 图题抽取：载荷 title > 引导语 > Mermaid title 指令 > 类型通用名
# ---------------------------------------------------------------------------

def _blocks(content: str):
    from app.routers.export import _parse_content_blocks
    return _parse_content_blocks(content)


def test_chart_json_block_uses_payload_title():
    content = ('脚手架专项施工管理组织架构如下图所示：\n\n'
               '```chart-json\n'
               '{"type":"architecture","title":"脚手架专项施工管理组织架构图",'
               '"root":{"label":"项目经理","children":[{"label":"技术负责人"}]}}\n'
               '```\n')
    charts = [b for b in _blocks(content) if b["type"] == "chart"]
    assert len(charts) == 1
    # 载荷 title 优先于引导语（AI 按本章语义撰写，最可信）
    assert charts[0]["title"] == "脚手架专项施工管理组织架构图"
    assert charts[0]["chart_type"] == "architecture"


def test_chart_json_block_falls_back_to_lead_in_title():
    content = ('悬挑层转换搭设顺序如下图所示：\n\n'
               '```chart-json\n'
               '{"type":"flowchart","steps":[{"id":"s1","label":"落地架搭设"},'
               '{"id":"s2","label":"悬挑梁安装"}],'
               '"edges":[{"from":"s1","to":"s2"}]}\n'
               '```\n')
    charts = [b for b in _blocks(content) if b["type"] == "chart"]
    assert len(charts) == 1
    assert charts[0]["title"] == "悬挑层转换搭设顺序"


def test_chart_json_block_without_type_is_inferred_not_defaulted_to_labor():
    """缺 type 键时按载荷结构推断 —— 旧实现静默兜底 labor → 必渲染失败。"""
    content = ('```chart-json\n'
               '{"title":"脚手架安全管理组织机构","root":{"label":"项目经理",'
               '"children":[{"label":"安全总监"}]}}\n'
               '```\n')
    charts = [b for b in _blocks(content) if b["type"] == "chart"]
    assert len(charts) == 1
    assert charts[0]["chart_type"] == "architecture"
    assert charts[0]["title"] == "脚手架安全管理组织机构"


def test_chart_json_block_uninferable_is_skipped():
    """结构也判不出类型 → 整块跳过（不占图号、不写红字占位）。"""
    content = '```chart-json\n{"title":"无结构载荷","note":"占位"}\n```\n'
    assert [b for b in _blocks(content) if b["type"] == "chart"] == []


def test_mermaid_block_uses_lead_in_title_over_directive():
    content = ('脚手架分段搭设与转换顺序如下图所示：\n\n'
               '```mermaid\n'
               'flowchart LR\n'
               ' A[基础验收] --> B[落地段搭设]\n'
               ' B --> C[悬挑段搭设]\n'
               '```\n')
    charts = [b for b in _blocks(content) if b["type"] == "chart"]
    assert len(charts) == 1
    assert charts[0]["title"] == "脚手架分段搭设与转换顺序"


def test_mermaid_block_without_lead_in_uses_title_directive():
    content = ('```mermaid\n'
               'gantt\n'
               '    title 施工进度计划\n'
               '    dateFormat X\n'
               '    施工准备 :a1, 0, 10d\n'
               '```\n')
    charts = [b for b in _blocks(content) if b["type"] == "chart"]
    assert len(charts) == 1
    assert charts[0]["title"] == "施工进度计划"


# ---------------------------------------------------------------------------
# ② 未闭合（截断）围栏：不渲染残片、不落代码块
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("lang,body", [
    ("mermaid", "flowchart LR\n A[基底清理] --> B[地基承载力检测]\n B --> C{"),
    ("chart-json", '{"type":"flowchart","steps":[{"id":"s1","label":"立杆定位"},'),
])
def test_unclosed_chart_fence_produces_no_block(lang, body):
    """围栏未闭合 = 生成被截断：残片既不渲染成图，也不作为代码块印进成稿。"""
    content = f"施工工艺流程如下图所示：\n\n```{lang}\n{body}\n"
    blocks = _blocks(content)
    assert [b for b in blocks if b["type"] in ("chart", "code")] == []


def test_unclosed_chart_fence_gives_back_following_prose():
    """截断残片之后的正文段落必须保留（旧实现被一并吞进代码块后静默丢失）。"""
    content = ("施工工艺流程如下图所示：\n\n"
               "```mermaid\nflowchart LR\n A[基底清理] --> B[地基检测]\n B --> C{\n"
               "放线完成后须经技术负责人复核，形成书面记录。\n")
    blocks = _blocks(content)
    texts = " ".join(b.get("text", "") for b in blocks if b["type"] == "paragraph")
    assert "放线完成后须经技术负责人复核" in texts, blocks
    assert [b for b in blocks if b["type"] in ("chart", "code")] == []


# ---------------------------------------------------------------------------
# ③ 结构推断器（chart_validators，登记侧/导出侧共用）
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("payload,expect", [
    ({"root": {"label": "项目经理"}}, "architecture"),
    ({"steps": [{"id": "s1", "label": "a"}]}, "flowchart"),        # steps 键即流程
    ({"zones": [{"id": "Z1", "name": "办公区"}]}, "layout"),
    ({"phases": ["准备", "主体"], "categories": ["普工"],
      "data": [[1], [2]]}, "labor"),
    ({"items": [{"label": "方案A", "value": 80}]}, "comparison"),
    ({"items": [{"name": "施工准备", "start": "2026-01-01", "end": "2026-01-20"}]}, "gantt"),
    ({"milestones": [{"name": "竣工验收", "day": 90}]}, "timeline"),
    ({"note": "无结构"}, ""),
    ("not-a-dict", ""),
    ({"type": "GANTT", "tasks": []}, "gantt"),      # 显式类型优先（大小写容错）
])
def test_infer_chart_type_from_payload(payload, expect):
    from app.services.chart_validators import infer_chart_type_from_payload
    assert infer_chart_type_from_payload(payload) == expect


# ---------------------------------------------------------------------------
# ④ 导出前内容体检（只体检不篡改）
# ---------------------------------------------------------------------------

def test_audit_content_counts_all_risk_classes():
    from app.routers.export import audit_content
    sections = [
        {"title": "工程规模与建筑分布",
         "content": "地下总建筑面积约××平方米，檐口高度约××m。"},
        {"title": "招标文件与技术标准要求",
         "content": "招标文件技术标对本项目提出要求，评标办法另见附件。"},
        {"title": "基础处理与立杆定位放线",
         "content": "<table><tr><td>参数</td></tr></table>\n"
                    "```mermaid\nflowchart LR\n A --> B\n B --> C{",
         },
        {"title": "计算书",
         "content": r"$\frac{}{}$ 与 \square 空参数"},
    ]
    audit = audit_content(sections)
    keys = {it["key"]: it for it in audit["items"]}
    assert keys["placeholder"]["count"] == 2
    assert keys["bidding_terms"]["count"] == 2
    assert keys["html_tag"]["count"] == 6          # table/tr/td 开闭标签各 1，按标签计
    assert keys["latex_square"]["count"] == 1
    assert keys["unclosed_fence"]["count"] == 1
    # 级别排序：error 在前（用户先看最要命的）
    levels = [it["level"] for it in audit["items"]]
    assert levels == sorted(levels, key=lambda x: 0 if x == "error" else 1)
    assert audit["total"] > 0


def test_audit_content_clean_document_has_no_items():
    from app.routers.export import audit_content
    audit = audit_content([{"title": "工程概况", "content": "本工程位于上海市，建筑檐口高度约98.5m。"}])
    assert audit == {"total": 0, "items": []}


def test_export_audit_headers_are_latin1_safe_and_skipped_when_clean():
    """响应头 Value 必须 latin-1 安全（中文直接写会抛 UnicodeEncodeError）。"""
    from app.routers.export import _export_audit_headers, audit_content
    assert _export_audit_headers({}) == {}
    assert _export_audit_headers({"total": 0, "items": []}) == {}

    audit = audit_content([{"title": "工程概况", "content": "约××平方米"}])
    headers = _export_audit_headers(audit)
    assert "X-Content-Audit" in headers
    raw = headers["X-Content-Audit"]
    raw.encode("latin-1")                      # 不抛异常即达标
    from urllib.parse import unquote
    restored = json.loads(unquote(raw))
    assert restored["items"][0]["key"] == "placeholder"
    assert restored["items"][0]["sections"] == ["工程概况"]


# ---------------------------------------------------------------------------
# ⑤ 端到端：导出图题必须用载荷真实标题（而非类型通用名）
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_export_docx_caption_uses_real_chart_title(db_conn):
    from app.routers.export import export_docx

    await db_conn.execute(
        "INSERT INTO schemes (id, project_id, name) VALUES (?,?,?)",
        ("sq", "pq", "脚手架专项施工方案"))
    content = (
        "悬挑层转换搭设顺序如下图所示：\n\n"
        "```chart-json\n"
        '{"type":"architecture","title":"悬挑层转换管理组织架构图",'
        '"root":{"label":"项目经理","children":[{"label":"技术负责人",'
        '"children":[{"label":"施工员"}]}]}}\n'
        "```\n"
    )
    await db_conn.execute(
        "INSERT INTO sections (id, scheme_id, parent_id, sort_order, level, title,"
        " status, content) VALUES (?,?,?,?,?,?,?,?)",
        ("sc1", "sq", "", 0, 1, "施工工艺技术", "generated", content))
    await db_conn.commit()

    resp = await export_docx("sq", {"config": {"allow_pil_fallback": True}}, db=db_conn)

    from docx import Document
    joined = "\n".join(p.text for p in Document(resp.path).paragraphs)
    assert "图 1-1 悬挑层转换管理组织架构图" in joined, joined
    # 旧行为：图题是类型通用名「组织架构图」，与本章"悬挑层转换"语义错配
    assert "图 1-1 组织架构图" not in joined, joined


# ---------------------------------------------------------------------------
# ⑥ 契约变更（v12）：渲染失败的图表默认跳过且不占用图号
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_export_docx_failed_chart_skipped_without_consuming_fig_num(db_conn):
    """默认跳过失败图：交付文档无红字报错，后续图的编号不虚跳（失败图不占号）。"""
    from app.routers.export import export_docx

    await db_conn.execute(
        "INSERT INTO schemes (id, project_id, name) VALUES (?,?,?)",
        ("sf", "pf", "脚手架专项施工方案-失败图"))
    content = (
        "劳动力配置如下图所示：\n\n"
        "```chart-json\n"        # 空壳 labor：校验型白名单内，但渲染必然失败
        '{"type":"labor"}\n'
        "```\n\n"
        "组织机构如下图所示：\n\n"
        "```chart-json\n"
        '{"type":"architecture","title":"悬挑层管理组织架构图",'
        '"root":{"label":"项目经理"}}\n'
        "```\n"
    )
    await db_conn.execute(
        "INSERT INTO sections (id, scheme_id, parent_id, sort_order, level, title,"
        " status, content) VALUES (?,?,?,?,?,?,?,?)",
        ("sf1", "sf", "", 0, 1, "施工工艺技术", "generated", content))
    await db_conn.commit()

    resp = await export_docx("sf", {"config": {"allow_pil_fallback": True}}, db=db_conn)

    from docx import Document
    joined = "\n".join(p.text for p in Document(resp.path).paragraphs)
    assert "渲染失败" not in joined, f"交付文档不应出现红字占位:\n{joined}"
    # 失败图不占号 → 成功的那张顶上是 1-1（旧行为会是 1-2）
    assert "图 1-1 悬挑层管理组织架构图" in joined, joined
    assert "图 1-2" not in joined, joined
    # 失败信息仍走系统侧信号（响应头）
    import json as _json
    stats = _json.loads(resp.headers["X-Chart-Render-Stats"])
    assert stats["failed"] >= 1, stats


@pytest.mark.asyncio
async def test_export_docx_failed_chart_placeholder_optin(db_conn):
    """chart_fail_placeholder=true 恢复 V7.0 占位形态（且图号照常占用）。"""
    from app.routers.export import export_docx

    await db_conn.execute(
        "INSERT INTO schemes (id, project_id, name) VALUES (?,?,?)",
        ("sp", "pp", "脚手架专项施工方案-占位"))
    content = (
        "劳动力配置如下图所示：\n\n"
        "```chart-json\n{\"type\":\"labor\"}\n```\n"
    )
    await db_conn.execute(
        "INSERT INTO sections (id, scheme_id, parent_id, sort_order, level, title,"
        " status, content) VALUES (?,?,?,?,?,?,?,?)",
        ("sp1", "sp", "", 0, 1, "施工工艺技术", "generated", content))
    await db_conn.commit()

    resp = await export_docx(
        "sp", {"config": {"allow_pil_fallback": True, "chart_fail_placeholder": True}},
        db=db_conn)

    from docx import Document
    joined = "\n".join(p.text for p in Document(resp.path).paragraphs)
    assert "渲染失败" in joined, f"占位模式应写入红字提示:\n{joined}"
    assert "图 1-1" in joined, joined


# ---------------------------------------------------------------------------
# ⑦ 剪贴板表格残留（v14）：<fcel>/<nl> 等标记不得以乱码形式印进成稿
# ---------------------------------------------------------------------------

def test_clean_text_strips_clipboard_table_markers():
    """Excel/Word 粘贴表格带入的占位标记与 HTML 表格标签降级为可读纯文本。"""
    from app.services.docx_math import clean_text
    # 纯标记堆叠（实测第 88 页乱码）→ 不残留任何尖括号标记
    garbage = "<br><br><table><fcel><fcel><nl></table>"
    out = clean_text(garbage)
    assert "<" not in out and ">" not in out, out
    # 带正文的 HTML 表格 → 保留单元格文字、去除包裹标签
    tbl = "<table><tr><td>参数A</td><td>值B</td></tr></table>"
    out2 = clean_text(tbl)
    assert "参数A" in out2 and "值B" in out2, out2
    assert "<table>" not in out2 and "<td>" not in out2, out2


def test_clean_text_keeps_legitimate_angle_brackets():
    """正文中合法的 < 、>（如「a < b」）不能被误删。"""
    from app.services.docx_math import clean_text
    assert clean_text("当 a < b 且 b > c 时").replace(" ", "") == "当a<b且b>c时"


def test_audit_content_counts_clipboard_table():
    from app.routers.export import audit_content
    audit = audit_content([{
        "title": "材料计划",
        "content": "下表为参数表：<br><table><fcel><fcel><nl></table>",
    }])
    keys = {it["key"]: it for it in audit["items"]}
    assert "clipboard_table" in keys
    # <fcel> ×2 + <nl> ×1 = 3
    assert keys["clipboard_table"]["count"] == 3


@pytest.mark.asyncio
async def test_export_docx_clipboard_table_not_printed_as_garbage(db_conn):
    """端到端：含剪贴板乱码的正文导出后成稿不再出现 <fcel>/<nl>/<table> 标记。"""
    from app.routers.export import export_docx

    await db_conn.execute(
        "INSERT INTO schemes (id, project_id, name) VALUES (?,?,?)",
        ("sc", "pc", "脚手架专项施工方案-乱码表"))
    await db_conn.execute(
        "INSERT INTO sections (id, scheme_id, parent_id, sort_order, level, title,"
        " status, content) VALUES (?,?,?,?,?,?,?,?)",
        ("sc1", "sc", "", 0, 1, "材料计划", "generated",
         "主要构配件如下：\n\n<br><br><table><fcel><fcel><nl></table>\n\n"
         "以上参数以计算书为准。"))
    await db_conn.commit()

    resp = await export_docx("sc", {"config": {}}, db=db_conn)

    from docx import Document
    joined = "\n".join(p.text for p in Document(resp.path).paragraphs)
    for noise in ("<fcel>", "<nl>", "<table>", "&lt;fcel&gt;", "&lt;nl&gt;"):
        assert noise not in joined, f"成稿不应残留表格乱码 {noise}:\n{joined}"
    # 正常正文不受影响
    assert "以上参数以计算书为准" in joined, joined


async def _seed_html_table_scheme(db_conn, scheme_id, section_id, content):
    await db_conn.execute(
        "INSERT INTO schemes (id, project_id, name) VALUES (?,?,?)",
        (scheme_id, "p" + scheme_id, "脚手架专项施工方案-HTML表"))
    await db_conn.execute(
        "INSERT INTO sections (id, scheme_id, parent_id, sort_order, level, title,"
        " status, content) VALUES (?,?,?,?,?,?,?,?)",
        (section_id, scheme_id, "", 0, 1, "搭设参数", "generated", content))
    await db_conn.commit()


_HTML_TABLE_CONTENT = (
    "搭设参数如下：\n\n"
    "<table><tr><th>项目</th><th>数值</th></tr>"
    "<tr><td>严禁使用 brick 架体</td><td>1.5m</td></tr></table>\n")


async def test_export_docx_auto_rewrite_html_table_becomes_native(db_conn):
    """auto_rewrite_content=true：HTML 表格转成原生 Word 表格，英文串写修正。"""
    from app.routers.export import export_docx
    await _seed_html_table_scheme(db_conn, "sca", "sca1", _HTML_TABLE_CONTENT)

    resp = await export_docx("sca", {"config": {"auto_rewrite_content": True}}, db=db_conn)

    from docx import Document
    doc = Document(resp.path)
    assert len(doc.tables) >= 1, "HTML 表格应被改写为原生 Word 表格"
    flat = [c.text.strip() for t in doc.tables for r in t.rows for c in r.cells]
    assert "严禁使用 砖料 架体" in flat, flat
    assert "1.5m" in flat, flat
    all_text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(flat)
    assert "brick" not in all_text and "<table>" not in all_text


async def test_export_docx_auto_rewrite_off_keeps_v14_behavior(db_conn):
    """默认（开关关）：行为与 v14 一致——表格扁平化无乱码，但不转原生表、英文不修正。"""
    from app.routers.export import export_docx
    await _seed_html_table_scheme(db_conn, "scb", "scb1", _HTML_TABLE_CONTENT)

    resp = await export_docx("scb", {"config": {}}, db=db_conn)

    from docx import Document
    doc = Document(resp.path)
    assert len(doc.tables) == 0, "未开启时不应将 HTML 表格转为原生表"
    joined = "\n".join(p.text for p in doc.paragraphs)
    assert "<table>" not in joined and "brick" in joined  # 扁平化保留原文，未修正英文
