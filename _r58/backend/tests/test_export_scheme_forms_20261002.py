"""第二十三轮（2026-10-02）：专项施工方案法定前置表单（导出）回归护栏。

本轮落地 L-0（AGENTS.md §4.22.9 未落地项 P1 最高项）：导出链路补齐
专项施工方案的法定交付形态。此前导出只有「封面 + 目录 + 正文」，
**编制说明 / 审批表 / 专家论证报告 / 施工图纸附件清单 四项零命中**。

规范依据（全部为部文原文，不凭记忆编写）：
  · 住建部令第37号 第十一条 —— 专项施工方案应当由施工单位技术负责人
    审核签字、加盖单位公章，并由总监理工程师审查签字、加盖执业印章
    后方可实施；分包的由总承包与分包技术负责人**共同**审核签字。
  · 住建部令第37号 第十二条 —— 专家从专家库选取，人数**不得少于 5 名**。
  · 住建部令第37号 第十三条 —— 专家论证会后形成**论证报告**，结论为
    **通过 / 修改后通过 / 不通过** 三选一，专家签字确认。
  · 建办质〔2018〕31号 第三条（参会人员五类）、第四条（论证内容三项）、
    二、第(九)项（计算书及相关施工图纸）。

红线：本表**不得出现**人名、证书编号、单位名称、日期等责任主体信息
（AGENTS.md 数据真实性红线）—— 表单的签字栏本就是「待签」空栏，
预填才是缺陷。

测试策略：配置规范化 + AST 静态断言 + 真实 DOCX 端到端渲染，零 AI。
"""
import ast
import io
import os
import sys

import pytest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

BACKEND = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
EX_PATH = os.path.join(BACKEND, "app", "routers", "export.py")

#: 四张表单的 key（与 _prepare_export.docx_options.scheme_forms 键名一致）
FORM_KEYS = ("compilation_note", "approval", "expert_review", "drawing_appendix")


def _src(path=EX_PATH):
    with io.open(path, encoding="utf-8") as f:
        return f.read()


def _fn(name, src=None):
    src = src or _src()
    for node in ast.parse(src).body:
        if isinstance(node, (ast.AsyncFunctionDef, ast.FunctionDef)) and node.name == name:
            return ast.get_source_segment(src, node) or ""
    raise AssertionError("未找到函数 " + name)


# ---------------------------------------------------------------------------
# 1. 开关默认关闭（向后兼容：产物与旧版逐字一致）
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("key", FORM_KEYS)
def test_scheme_forms_config_defaults_off(key):
    """未配置 / 全关 → scheme_forms 为空 dict → _build_docx_sync 不渲染任何表单。"""
    from app.routers.export import _prepare_export  # 触发导入，确认可加载
    assert _prepare_export is not None
    # 规范化逻辑：空输入与脏输入都得空 dict
    src = _fn("_prepare_export")
    i = src.find('"scheme_forms"')
    assert i > 0, "_prepare_export 未读取 scheme_forms"
    seg = src[i:i + 700]
    assert "bool(" in seg, "scheme_forms 未做 bool 化（脏值会渗进渲染分支）"
    # 四个键全部登记
    for k in FORM_KEYS:
        assert f'"{k}"' in seg, f"未登记表单开关 {k}"


def test_scheme_forms_in_fingerprint_whitelist():
    """必须进 _EXPORT_CONFIG_KEYS：否则切换开关后 config_hash 不变 →
    命中旧缓存 → 用户打开开关却看不到表单。"""
    src = _src()
    i = src.find("_EXPORT_CONFIG_KEYS = frozenset({")
    assert i > 0
    seg = src[i:src.find("})", i)]
    assert '"scheme_forms"' in seg, "scheme_forms 未纳入内容指纹白名单"


def test_render_branch_guarded_by_truthiness():
    """渲染分支必须以「非空 dict」为门禁，默认不进入。"""
    src = _fn("_build_docx_sync")
    assert "scheme_forms" in src
    assert "if _forms:" in src, "渲染分支必须以 _forms 非空为门禁"
    # 四张表都受开关控制（不得有绕过分支）
    for k in FORM_KEYS:
        assert f'"{k}"' in src, f"表单 {k} 未接入渲染分支"


# ---------------------------------------------------------------------------
# 2. 规范依据：条文不得被改写
# ---------------------------------------------------------------------------

def test_approval_table_cites_decree37_article11():
    """审批表依据必须是住建部令第37号第十一条（签字方与盖章方）。"""
    src = _fn("_add_scheme_approval_table")
    assert "第十一条" in src, "审批表未标注依据条文"
    assert "37号" in src
    for who in ("施工单位技术负责人", "总监理工程师"):
        assert who in src, f"审批表缺法定签字方：{who}"
    # 分包时共同审核（第十一条第二款）
    assert "总承包单位技术负责人" in src and "专业分包单位技术负责人" in src, \
        "审批表缺分包工程共同审核签字栏（37号令第十一条第二款）"
    # 执业印章 / 单位公章
    assert "执业印章" in src and "单位公章" in src


def test_expert_review_cites_decree37_article12_13():
    """专家论证报告：人数不得少于 5 名 + 结论三选一（37号令第十二/十三条）。"""
    src = _fn("_add_expert_review_form")
    assert "第十二条" in src and "第十三条" in src
    assert "不得少于 5 名" in src, "未写入「专家人数不得少于 5 名」（第十二条）"
    # 结论三选一，逐字
    for opt in ("通过", "修改后通过", "不通过"):
        assert opt in src, f"论证结论缺选项：{opt}"
    # 固定 5 行专家签字空栏
    assert "range(1, 6)" in src, "专家签字栏未按 5 名生成"
    # 31号文 论证内容三项，逐字
    for item in ("是否完整、可行", "计算书和验算依据、施工图",
                 "是否满足现场实际情况"):
        assert item in src, f"论证内容缺项（31号文第四条）：{item}"
    # 31号文 参会人员五类
    for who in ("建设单位项目负责人", "勘察、设计单位",
                "专职安全生产管理人员", "总监理工程师"):
        assert who in src, f"参会人员缺类（31号文第三条）：{who}"


def test_drawing_appendix_cites_clause9():
    """施工图纸附件清单依据建办质〔2018〕31号 第(九)项。"""
    src = _fn("_add_drawing_appendix_page")
    assert "第（九）项" in src or "第(九)项" in src
    assert "31号" in src
    assert "图号" in src and "图纸名称" in src


def test_compilation_note_mentions_review_chain():
    """编制说明须写明「审核签字 / 审查签字后方可实施」（37号令第十一条）。"""
    src = _fn("_add_compilation_note_page")
    assert "审核签字" in src and "审查签字" in src
    assert "方可实施" in src


# ---------------------------------------------------------------------------
# 3. 红线：不得预填责任主体信息
# ---------------------------------------------------------------------------

def test_forms_never_prefill_signatures():
    """签字栏必须留空 —— 预填人名/证书号是「编造责任主体信息」。

    ⚠️ 判据锚点：只查**交付文档的渲染结果**与签字行字面量，不做源码全文
    关键词扫描。源码里本函数自己的红线注释就写着「不编造人名、证书编号」，
    全文匹配必然被自己的说明文字误伤 —— AGENTS.md §5.14 记录的
    「锚点选宽了会因注释含关键词而恒失败，逼着后人把护栏改松」同构陷阱。
    要拦的是**产物里的编造**，就锚产物本身。
    """
    for fn_name, labels in (
        ("_add_scheme_approval_table",
         ("编制人（签字）", "总监理工程师（审查签字）",
          "施工单位技术负责人（审核签字）")),
        # 专家签字栏由 f-string 生成（`f"专家 {i}（签字）"`），源码里不是
        # 「专家 」字面量，只能断言生成表达式本身带空值。
        ("_add_expert_review_form", ('f"专家 {i}（签字）", ""',)),
    ):
        src = _fn(fn_name)
        for label in labels:
            i = src.find(label)
            assert i > 0, f"{fn_name} 缺签字栏 {label}"
            line = src[i:src.find("\n", i)]
            assert '""' in line, f"{fn_name} 的 {label} 被预填了值：{line.strip()}"


def test_rendered_document_has_no_fabricated_identity(tmp_path):
    """产物中不得出现责任主体标识（人名 / 证书编号形态）。"""
    out = _build(tmp_path, forms={k: True for k in FORM_KEYS},
                 name="noid.docx")
    text = _text(out)
    # 证书编号形态：字母数字串 + 「证号/证书编号」；中文姓名形态：3 字全中文
    import re
    assert not re.search(r"[0-9A-Z]{6,}\s*(?:证号|证书编号)", text), \
        "产物出现证书编号形态（编造责任主体信息）"
    assert "证书编号" not in text, "产物出现「证书编号」字样"
    # 签字栏确实存在（否则「没有编造」是因为根本没渲染出来）
    assert "编制人（签字）" in text, "签字栏缺失，无法证明未预填"


def test_forms_filter_empty_header_rows(tmp_path):
    """抬头行空值必须跳过（交付文档出现「工程名称：」空栏比缺行更糟）。"""
    out = _build(tmp_path, forms={"approval": True}, name="emptyhdr.docx")
    text = _text(out)
    # 未传 cover_info → 不应出现空的「工程名称」「方案编号」抬头行
    assert "工程名称" not in text, "空抬头行未被过滤"
    assert "方案编号" not in text, "空抬头行未被过滤"
    # 抬头行的过滤不得牵连签字栏（签字栏是空值但必须保留）
    assert "编制人（签字）" in text, "签字栏被空值过滤误伤"


def test_form_render_fail_soft():
    """前置表单渲染异常只记 WARNING，不阻断正文导出。"""
    src = _fn("_build_docx_sync")
    i = src.find("if _forms:")
    assert i > 0
    seg = src[i:i + 3000]
    assert "except Exception" in seg, "表单渲染未做 fail-soft 兜底"
    assert "logger.warning" in seg, "表单渲染失败必须留日志痕迹"


# ---------------------------------------------------------------------------
# 4. 端到端：真实 DOCX 渲染
# ---------------------------------------------------------------------------

def _build(tmp_path, *, forms=None, cover_info=None, name="formtest.docx"):
    from app.routers.export import _build_docx_sync, _load_heading_styles, _parse_content_blocks
    content = "本章为施工工艺技术内容，工艺参数按专项计算书核定。"
    sections = [{
        "id": "c1", "parent_id": "", "sort_order": 0, "level": 1,
        "title": "第一章 施工工艺技术", "content": content,
    }]
    blocks = {"c1": _parse_content_blocks(content)}
    out = str(tmp_path / name)
    _build_docx_sync(
        out, {"id": "s1", "project_id": "p1", "name": "深基坑开挖专项施工方案"},
        sections, {}, {}, {}, blocks,
        "宋体", 12, "", "", False, False, False,
        bidder_name="",
        heading_styles=_load_heading_styles({}),
        page_break_before_chapter=True, line_spacing=1.5,
        page_number_style="simple", toc_depth=3,
        cover_info=cover_info,
        scheme_forms=forms or {},
    )
    return out


def _text(path):
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    return "\n".join(parts)


def test_all_four_forms_render(tmp_path):
    """四张表全部开启 → 页题与法定要素都出现在交付文档里。"""
    out = _build(tmp_path, forms={k: True for k in FORM_KEYS},
                 cover_info={"工程名称": "示例工程", "方案编号": "SF-2026-001",
                             "编制单位": "示例施工单位"})
    text = _text(out)
    assert "编制说明" in text
    assert "专项施工方案审批表" in text
    assert "专项施工方案专家论证报告" in text
    assert "施工图纸附件清单" in text
    # 抬头带出封面信息
    assert "示例工程" in text and "SF-2026-001" in text
    # 论证内容三项
    assert "计算书和验算依据、施工图" in text
    # 专家 5 名空栏
    assert text.count("专家 ") >= 5, "专家签字栏不足 5 名"
    # 表格确实生成了（_add_form_table 产出 Word 表格）
    from docx import Document
    assert len(Document(out).tables) >= 5, "表单表格未生成"


def test_forms_absent_by_default(tmp_path):
    """默认关闭：四张表都不出现（向后兼容）。"""
    out = _build(tmp_path)
    text = _text(out)
    for token in ("编制说明", "专项施工方案审批表",
                  "专项施工方案专家论证报告", "施工图纸附件清单"):
        assert token not in text, f"默认关闭时不应出现 {token}"


def test_forms_individually_switchable(tmp_path):
    """逐项独立：只开审批表时，审批表在、其余三张不在。"""
    out = _build(tmp_path, forms={"approval": True}, name="only_approval.docx")
    text = _text(out)
    assert "专项施工方案审批表" in text
    assert "编制说明" not in text
    assert "专项施工方案专家论证报告" not in text
    assert "施工图纸附件清单" not in text


def test_forms_all_false_equals_absent(tmp_path):
    """全 False 与不传等价（不留空白页）。"""
    out = _build(tmp_path, forms={k: False for k in FORM_KEYS}, name="allfalse.docx")
    text = _text(out)
    assert "专项施工方案审批表" not in text
    from docx import Document
    doc = Document(out)
    # 无表单 → 表格数为 0（正文无表格）
    assert len(doc.tables) == 0, "全关时不应生成任何表单表格"


def test_export_module_parses():
    ast.parse(_src(), filename=EX_PATH)
