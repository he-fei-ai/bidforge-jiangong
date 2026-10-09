"""导出层「列表样式被拍平」修复回归测试（2026-10-03）

背景：write_section 的有序列表序号原用单一全局计数器（ordered_seq /
ordered_marker），导致嵌套有序列表的子项被拍平成一条贯穿所有层级的连续
序列（"1. 顶层 / 2. 子项 / 3. 子项 / 4. 顶层二"），子项未在父级内重新从 1
计数。修复改为按缩进层级（indent_lvl）的分层栈 ordered_stack，每层独立维护
(marker, seq)。

本文件锁定：
- 嵌套有序列表按层级重计数 + 父级序号续接（修复核心）；
- 嵌套有序 + 不同中文枚举符混合；
- 三层级嵌套；
- 深层跳级（indent 不连续）不崩；
- 既有契约全部保留：无序子项不重置有序序号、遇非列表块重置序列、
  标记样式变化本层从 1 重计、缩进层级缩进保留。
- 反向锁定：旧的「拍平连续序列」形态不得出现。
"""
import os
import tempfile

from app.routers.export import (
    _build_docx_sync,
    _load_heading_styles,
    _parse_content_blocks,
)
from docx import Document


def _build_single(content: str) -> list[str]:
    """复用导出单章节入口，返回正文段落文本（去掉章节标题）。"""
    sections = [{
        "id": "c1", "parent_id": "", "sort_order": 0, "level": 1,
        "title": "第一章 测试章节", "content": content,
    }]
    blocks = {"c1": _parse_content_blocks(content)}
    out = os.path.join(tempfile.gettempdir(), "nest_ordered.docx")
    _build_docx_sync(
        out, {"id": "s1", "project_id": "p1", "name": "测试方案"},
        sections, {}, {}, {}, blocks,
        "宋体", 12, "", "", False, False, False,
        bidder_name="",
        heading_styles=_load_heading_styles({}),
        page_break_before_chapter=True,
        line_spacing=1.5,
        page_number_style="simple",
        toc_depth=3,
        margins=None, cover_info=None,
    )
    return [p.text for p in Document(out).paragraphs
            if p.text.strip() and p.text.strip() != "第一章 测试章节"]


class TestNestedOrderedRenumbersPerLevel:
    """✅ 修复核心：嵌套有序子项在父级内重新从 1 计数，父级序号返回时续接。"""

    def test_two_level_nesting(self):
        paras = _build_single(
            "1. 顶层一\n"
            "   1. 子项一\n"
            "   2. 子项二\n"
            "2. 顶层二\n"
        )
        assert "1. 顶层一" in paras
        assert "1. 子项一" in paras
        assert "2. 子项二" in paras
        assert "2. 顶层二" in paras

    def test_flattened_form_absent(self):
        """反向锁定：旧的拍平连续序列（子项被当成 2./3.）不得出现。"""
        paras = _build_single(
            "1. 顶层一\n"
            "   1. 子项一\n"
            "   2. 子项二\n"
            "2. 顶层二\n"
        )
        assert not any(t.startswith("2. 子项一") for t in paras), paras
        assert not any(t.startswith("3. 子项二") for t in paras), paras
        assert not any(t.startswith("4. 顶层二") for t in paras), paras

    def test_three_level_nesting(self):
        paras = _build_single(
            "1. A\n"
            "   1. B\n"
            "      1. C\n"
            "   2. B2\n"
            "2. A2\n"
        )
        assert "1. A" in paras
        assert "1. B" in paras
        assert "1. C" in paras
        assert "2. B2" in paras
        assert "2. A2" in paras
        # 反向：C 不应被排成 3.
        assert not any(t.startswith("3. C") for t in paras), paras


class TestNestedOrderedMixedMarker:
    """嵌套有序 + 中文枚举符混合：每层按自己 marker 渲染。"""

    def test_ordered_parent_cn_child(self):
        paras = _build_single(
            "1. 顶层一\n"
            "   （一）子甲\n"
            "   （二）子乙\n"
            "2. 顶层二\n"
        )
        assert "1. 顶层一" in paras
        assert "（一）子甲" in paras
        assert "（二）子乙" in paras
        assert "2. 顶层二" in paras

    def test_cn_parent_ordered_child(self):
        paras = _build_single(
            "（一）设计标准\n"
            "   1. 子点一\n"
            "   2. 子点二\n"
            "（二）材料要求\n"
        )
        assert "（一）设计标准" in paras
        assert "1. 子点一" in paras
        assert "2. 子点二" in paras
        assert "（二）材料要求" in paras


class TestLegacyContractsPreserved:
    """既有契约在修复后必须全部保留（不得回归）。"""

    def test_unordered_sibling_keeps_ordered_seq(self):
        """无序子项不碰栈 → 父级有序序号续接（非重置）。"""
        paras = _build_single(
            "1. 第一步：测量放线\n"
            "- 注意事项：避开雨天\n"
            "2. 第二步：基坑开挖\n"
        )
        assert "1. 第一步：测量放线" in paras
        assert "• 注意事项：避开雨天" in paras
        assert "2. 第二步：基坑开挖" in paras
        assert not any(t.startswith("1. 第二步") for t in paras), paras

    def test_non_list_block_resets_ordered_seq(self):
        """遇非列表块（过渡段落）→ 重置整条有序序列。"""
        paras = _build_single(
            "（一）设计标准\n"
            "（二）材料要求\n"
            "\n"
            "正文过渡段落。\n"
            "\n"
            "1、施工准备\n"
            "1、测量放线\n"
        )
        assert "（一）设计标准" in paras
        assert "（二）材料要求" in paras
        assert "1、施工准备" in paras
        assert "2、测量放线" in paras  # 重复编号自动连续

    def test_marker_change_resets_same_level(self):
        """同层标记样式变化 → 该层从 1 重计（不沿用 ascii 序号）。"""
        paras = _build_single(
            "1. A\n"
            "2. B\n"
            "（一）C\n"
            "（二）D\n"
        )
        assert "1. A" in paras
        assert "2. B" in paras
        assert "（一）C" in paras
        assert "（二）D" in paras
        assert not any(t.startswith("3. C") for t in paras), paras

    def test_unordered_nested_indent_preserved(self):
        """无序嵌套缩进层级保留（0.74 / 1.48 cm），未被拍平到同一级。"""
        from docx import Document as _D
        sections = [{
            "id": "c1", "parent_id": "", "sort_order": 0, "level": 1,
            "title": "第一章 测试章节",
            "content": "- 顶层A\n  - 二级A1\n  - 二级A2\n- 顶层B\n",
        }]
        content = sections[0]["content"]
        blocks = {"c1": _parse_content_blocks(content)}
        out = os.path.join(tempfile.gettempdir(), "nest_unordered.docx")
        _build_docx_sync(
            out, {"id": "s1", "project_id": "p1", "name": "t"},
            sections, {}, {}, {}, blocks,
            "宋体", 12, "", "", False, False, False,
            bidder_name="", heading_styles=_load_heading_styles({}),
            page_break_before_chapter=True, line_spacing=1.5,
            page_number_style="simple", toc_depth=3,
            margins=None, cover_info=None,
        )
        indents = []
        for p in _D(out).paragraphs:
            if p.text.strip() and p.text.strip() != "第一章 测试章节":
                li = p.paragraph_format.left_indent
                indents.append(round(li.cm, 2) if li else 0.0)
        # 顶层 0.74，二层 1.48 —— 两级不同，证明未被拍平
        assert indents.count(0.74) == 2, indents
        assert indents.count(1.48) == 2, indents


class TestDeepLevelSkipNoCrash:
    """indent 不连续（跳级）不得导致越界/崩溃。"""

    def test_skip_two_levels(self):
        paras = _build_single(
            "1. A\n"
            "      1. deep\n"   # 6 空格 -> level 3（中间层级被补齐）
            "2. A2\n"
        )
        assert "1. A" in paras
        assert "1. deep" in paras
        assert "2. A2" in paras
