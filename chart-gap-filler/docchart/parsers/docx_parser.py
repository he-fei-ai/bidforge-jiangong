# -*- coding: utf-8 -*-
"""DOCX 解析器：按 body 顺序遍历段落与表格。

每个块保留原生对象引用（native）与 body 序号（para_idx），
导出时直接在原生位置之后回插图片段落，最大限度保留原格式。
"""

from __future__ import annotations

from pathlib import Path

from ..models import Block, BlockKind, Document, figure_caption_of

_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_DRAWING_NS = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"


def _para_style(para) -> str:
    try:
        return (para.style.name or "").lower()
    except Exception:
        return ""


def _has_image(para) -> bool:
    """段落是否包含图片/图形（w:drawing 或旧式 pict）。"""
    xml = para._p.xml
    return "<w:drawing" in xml or "<w:pict" in xml or "v:imagedata" in xml


def _is_heading(para) -> tuple[bool, int]:
    st = _para_style(para)
    if st.startswith("heading"):
        try:
            return True, int(st.split()[-1])
        except (ValueError, IndexError):
            return True, 1
    if st == "title":
        return True, 1
    # 无样式时按"第X章/编号开头且很短"粗判（针对国内方案文档常见直接格式）
    text = para.text.strip()
    if text and len(text) <= 40:
        import re
        if re.match(r"^(第[一二三四五六七八九十百\d]+章|[1-9]\d?\s)", text):
            lvl = 1 if text.startswith("第") else 2
            return True, lvl
    return False, 0


def parse_docx(path: str | Path) -> Document:
    from docx import Document as DocxDocument  # 延迟导入，避免非 docx 场景硬依赖

    p = Path(path)
    doc = DocxDocument(str(p))
    blocks: list[Block] = []
    section_stack: list[tuple[int, str]] = []

    body = doc.element.body
    para_iter = iter(doc.paragraphs)
    tbl_iter = iter(doc.tables)
    cur_para, cur_tbl = next(para_iter, None), next(tbl_iter, None)
    idx = 0

    def _section() -> str:
        return " > ".join(t for _, t in section_stack)

    def _push(**kw) -> None:
        kw.setdefault("section", _section())
        kw["id"] = len(blocks)
        blocks.append(Block(**kw))

    for child in body:
        tag = child.tag
        if tag == f"{_W_NS}p" and cur_para is not None:
            para = cur_para
            cur_para = next(para_iter, None)
            text = para.text.strip()
            style = _para_style(para)
            is_img = _has_image(para)
            if is_img and not text:
                _push(kind=BlockKind.IMAGE, text="[内嵌图片]",
                      src={"para_idx": idx}, native=para)
            elif is_img:
                # 图文混排：既算图片也保留文字
                _push(kind=BlockKind.IMAGE, text=text,
                      src={"para_idx": idx}, native=para)
            elif "caption" in style or "题注" in style:
                _push(kind=BlockKind.CAPTION, text=text,
                      src={"para_idx": idx}, native=para)
            elif figure_caption_of(text):
                _push(kind=BlockKind.CAPTION, text=text,
                      src={"para_idx": idx}, native=para)
            else:
                hd, level = _is_heading(para)
                if hd:
                    while section_stack and section_stack[-1][0] >= level:
                        section_stack.pop()
                    section_stack.append((level, text))
                    _push(kind=BlockKind.HEADING, text=text, level=level,
                          src={"para_idx": idx}, native=para)
                elif "list" in style:
                    _push(kind=BlockKind.LIST, text=text,
                          src={"para_idx": idx}, native=para)
                elif text:
                    _push(kind=BlockKind.PARAGRAPH, text=text,
                          src={"para_idx": idx}, native=para)
                else:
                    # 空段落也登记（锚点回插时可能用到）
                    _push(kind=BlockKind.PARAGRAPH, text="",
                          src={"para_idx": idx}, native=para)
            idx += 1
        elif tag == f"{_W_NS}tbl" and cur_tbl is not None:
            table = cur_tbl
            cur_tbl = next(tbl_iter, None)
            rows = []
            for row in table.rows:
                cells = []
                seen = set()
                for c in row.cells:  # 合并单元格会重复引用同一 tc，去重
                    if id(c._tc) in seen:
                        continue
                    seen.add(id(c._tc))
                    cells.append(" ".join(c.text.split()))
                rows.append(cells)
            # 表注（"表 x-x 标题"）识别为前一段落，这里只登记表格本体
            _push(kind=BlockKind.TABLE,
                  text=" | ".join(rows[0]) if rows else "",
                  rows=rows, src={"para_idx": idx}, native=table)
            idx += 1

    # docx 的 raw 用段落全文拼接，供报告摘要引用；native 保留原对象供回插
    raw = "\n".join(b.text for b in blocks if b.text)
    return Document(source_path=str(p), fmt="docx", blocks=blocks, raw=raw,
                    native=doc)
