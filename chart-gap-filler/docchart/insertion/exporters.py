# -*- coding: utf-8 -*-
"""导出器：把已生成图片按锚点回插到文档副本，输出 Markdown/HTML/DOCX。

共同约定：
- 不修改原文档，一律写出新文件（默认 *_autofill 后缀）；
- Markdown/HTML 按源位置（行号/字符偏移）从后往前回插，原格式零重排；
- DOCX 直接在原生段落元素前后插入新段落，保留全部原有内容与样式。
"""

from __future__ import annotations

import logging
import os
from pathlib import Path

from ..config import Config
from ..models import ChartGap, Document

logger = logging.getLogger("docchart.export")


def _rel_img(out_path: Path, img_path: str) -> str:
    """计算导出文档中引用图片的相对路径（posix 风格）。

    无法相对化时回退绝对路径：Windows 跨盘符 relpath 会抛 ValueError，
    不能让它炸掉整条导出链。
    """
    try:
        rel = os.path.relpath(img_path, out_path.parent)
    except ValueError:
        rel = str(Path(img_path).absolute())
    return rel.replace("\\", "/")


def _caption_text(g: ChartGap, cfg: Config) -> str:
    """图注正文：检测阶段已带补全标记时不重复追加后缀。"""
    suffix = str(cfg.get("output", "auto_caption_suffix", ""))
    return g.caption if suffix and suffix in g.caption else f"{g.caption}{suffix}"


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------

def export_markdown(doc: Document, gaps: list[ChartGap], out_path: Path,
                    cfg: Config) -> Path:
    """在锚点行之后（或图注行之前）插入图片与图注。"""
    lines = doc.raw.splitlines(keepends=True)
    blocks = {b.id: b for b in doc.blocks}
    inserts: list[tuple[int, str]] = []  # (行号, 文本)
    for g in [x for x in gaps if x.status == "filled"]:
        b = blocks[g.anchor_block_id]
        at = b.src["line_start"] if g.insert_before else b.src["line_end"] + 1
        rel = _rel_img(out_path, g.output_image)
        snippet = (f"\n![{g.caption}]({rel})\n\n"
                   f"{_caption_text(g, cfg)}\n\n")
        inserts.append((at, snippet))
    # 从后往前插，避免行号漂移
    out_lines = lines[:]
    for at, snippet in sorted(inserts, key=lambda x: -x[0]):
        at = min(max(at, 0), len(out_lines))
        out_lines.insert(at, snippet)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text("".join(out_lines), encoding="utf-8")
    return out_path


# ---------------------------------------------------------------------------
# HTML
# ---------------------------------------------------------------------------

def export_html(doc: Document, gaps: list[ChartGap], out_path: Path,
                cfg: Config) -> Path:
    raw = doc.raw
    blocks = {b.id: b for b in doc.blocks}
    inserts: list[tuple[int, str]] = []
    for g in [x for x in gaps if x.status == "filled"]:
        b = blocks[g.anchor_block_id]
        pos = b.src["html_start"] if g.insert_before else b.src["html_end"]
        rel = _rel_img(out_path, g.output_image)
        snippet = (f'\n<figure class="auto-filled"><img src="{rel}" alt="{g.caption}">'
                   f'<figcaption>{_caption_text(g, cfg)}</figcaption></figure>\n')
        inserts.append((pos, snippet))
    for pos, snippet in sorted(inserts, key=lambda x: -x[0]):
        raw = raw[:pos] + snippet + raw[pos:]
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(raw, encoding="utf-8")
    return out_path


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def export_docx(doc: Document, gaps: list[ChartGap], out_path: Path,
                cfg: Config) -> Path:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.shared import Cm
    from docx.text.paragraph import Paragraph

    native = doc.native
    if native is None:
        raise RuntimeError("DOCX 原生对象丢失，无法回插")
    blocks = {b.id: b for b in doc.blocks}
    filled = [x for x in gaps if x.status == "filled"]

    for g in filled:
        b = blocks[g.anchor_block_id]
        if Path(g.output_image).suffix.lower() == ".svg":
            # python-docx 不支持嵌入 SVG（Word 需 png 回退图），跳过回插但保留图片文件
            logger.warning("DOCX 无法嵌入 SVG，跳过回插：%s", g.output_image)
            continue
        anchor_el = getattr(b.native, "_p", None)
        if anchor_el is None:
            anchor_el = getattr(b.native, "_tbl", None)
        if anchor_el is None:
            logger.warning("锚点块 %d 无原生元素，跳过回插", b.id)
            continue
        # 先建图段落，再建图注段落（同一父容器下 addnext/addprevious 顺序稳定）
        img_p = OxmlElement("w:p")
        cap_p = OxmlElement("w:p")
        if g.insert_before:
            anchor_el.addprevious(img_p)
            img_p.addnext(cap_p)
        else:
            anchor_el.addnext(cap_p)
            anchor_el.addnext(img_p)
        img_para = Paragraph(img_p, b.native._parent if hasattr(b.native, "_parent") else native)
        run = img_para.add_run()
        run.add_picture(g.output_image, width=Cm(14.6))
        img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_para = Paragraph(cap_p, native)
        cap_para.add_run(_caption_text(g, cfg))
        cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    out_path.parent.mkdir(parents=True, exist_ok=True)
    native.save(str(out_path))
    return out_path


# ---------------------------------------------------------------------------
# 注册表入口
# ---------------------------------------------------------------------------

_EXPORTERS = {
    "markdown": export_markdown,
    "html": export_html,
    "docx": export_docx,
}


def register_exporter(fmt: str, fn):
    """插件扩展点：注册新的导出器。"""
    _EXPORTERS[fmt] = fn


def export_document(doc: Document, gaps: list[ChartGap], out_path: str | Path,
                    cfg: Config) -> Path:
    out_path = Path(out_path)
    fn = _EXPORTERS.get(doc.fmt)
    if fn is None:
        raise ValueError(f"未注册 {doc.fmt} 的导出器")
    p = fn(doc, gaps, out_path, cfg)
    logger.info("导出完成：%s（回插 %d 图）", p,
                sum(1 for g in gaps if g.status == "filled"))
    return p
