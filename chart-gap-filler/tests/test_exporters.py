# -*- coding: utf-8 -*-
"""导出器组件级交互测试：锚点回插、图注拼接、格式守卫。

覆盖 exporters 与 pipeline/models 的数据传递契约：
- Markdown 按行号、HTML 按字符偏移、DOCX 按原生元素回插；
- 图注后缀去重、相对路径 posix 化、svg 嵌入守卫、未注册格式报错。
"""

from pathlib import Path

import pytest

from docchart.config import Config
from docchart.insertion.exporters import (_caption_text, _rel_img,
                                          export_document, export_docx,
                                          export_html, export_markdown)
from docchart.models import ChartGap, Document, GapType
from docchart.parsers import parse_file

SUFFIX = "（AI 补全）"


def _gap(anchor: int, img: str = "charts/a.png", caption: str = "图 1-1 测试",
         insert_before: bool = False, status: str = "filled") -> ChartGap:
    return ChartGap(anchor_block_id=anchor, gap_type=GapType.EMPTY_PLACEHOLDER,
                    reason="", confidence=1.0, suggested_type="bar",
                    caption=caption, insert_before=insert_before,
                    status=status, output_image=img)


def _md(tmp_path: Path, text: str) -> Document:
    src = tmp_path / "in.md"
    src.write_text(text, encoding="utf-8")
    return parse_file(src)


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------

def test_markdown_insert_after_anchor(tmp_path: Path):
    doc = _md(tmp_path, "# 标题\n\n正文段落。\n")
    out = tmp_path / "o.md"
    export_markdown(doc, [_gap(1)], out, Config())
    text = out.read_text(encoding="utf-8")
    # 图片与图注都在锚点段落之后，且原文完整保留
    assert text.index("正文段落。") < text.index("![") < text.index(f"图 1-1 测试{SUFFIX}")


def test_markdown_insert_before_anchor(tmp_path: Path):
    doc = _md(tmp_path, "# 标题\n\n图 1-1 测试\n")
    out = tmp_path / "o.md"
    export_markdown(doc, [_gap(1, insert_before=True)], out, Config())
    text = out.read_text(encoding="utf-8")
    # 回插图位于原图注行（无后缀、独占一行）之前；原图注保留
    assert text.index("![") < text.index("图 1-1 测试\n")
    assert SUFFIX in text and text.count("# 标题") == 1


def test_markdown_relative_image_path_is_posix(tmp_path: Path):
    img = str(tmp_path / "charts" / "a.png")
    assert "\\" not in _rel_img(tmp_path / "out" / "x.md", img)


def test_rel_img_never_raises_on_unrelativizable_path():
    # 跨盘符路径无法相对化，必须回退绝对路径而非炸掉导出链
    r = _rel_img(Path("C:/docs/o.md"), "J:/assets/img/a.png")
    assert r.replace("\\", "/").endswith("a.png") and "\\" not in r


# ---------------------------------------------------------------------------
# HTML
# ---------------------------------------------------------------------------

def test_html_multiple_inserts_keep_original_offsets(tmp_path: Path):
    src_html = ("<html><body><h1>标题</h1><p>第一段。</p>"
                "<p>第二段带【插入图表】。</p></body></html>")
    src = tmp_path / "in.html"
    src.write_text(src_html, encoding="utf-8")
    doc = parse_file(src)
    out = tmp_path / "o.html"
    gaps = [_gap(1, img="c/1.png"), _gap(2, img="c/2.png", caption="图 2-1 乙")]
    export_html(doc, gaps, out, Config())
    text = out.read_text(encoding="utf-8")
    assert text.count('<figure class="auto-filled">') == 2
    # 原始块顺序与内容不因插入而漂移
    assert (text.index("<h1>标题</h1>") < text.index("第一段。")
            < text.index("第二段带【插入图表】。"))


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def test_docx_skips_svg_without_crash(tmp_path: Path):
    from docx import Document as Docx
    src = tmp_path / "in.docx"
    d = Docx()
    d.add_paragraph("图 1-1 节点示意")
    d.save(src)
    doc = parse_file(src)
    out = tmp_path / "o.docx"
    g = _gap(doc.blocks[0].id, img=str(tmp_path / "c.svg"))
    export_docx(doc, [g], out, Config())          # 不应抛异常
    assert len(Docx(str(out)).inline_shapes) == 0  # svg 被跳过而非写入失败


def test_docx_anchor_table_inserts_picture(tmp_path: Path):
    from docx import Document as Docx
    src = tmp_path / "in.docx"
    d = Docx()
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "月"
    t.cell(0, 1).text = "量"
    t.cell(1, 0).text = "3月"
    t.cell(1, 1).text = "10"
    d.add_paragraph("尾段")
    d.save(src)
    doc = parse_file(src)
    tbl = next(b for b in doc.blocks if b.kind.value == "table")
    png = tmp_path / "c.png"
    _make_png(png)
    out = tmp_path / "o.docx"
    export_docx(doc, [_gap(tbl.id, img=str(png))], out, Config())
    d2 = Docx(str(out))
    assert len(d2.inline_shapes) == 1
    # 图段落插在表格之后、尾段之前
    texts = [p.text for p in d2.paragraphs if p.text]
    assert texts[-1] == "尾段"


def _make_png(path: Path) -> None:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig = plt.figure()
    fig.savefig(str(path))
    plt.close(fig)


# ---------------------------------------------------------------------------
# 图注与注册表
# ---------------------------------------------------------------------------

def test_caption_suffix_not_duplicated():
    cfg = Config()
    assert _caption_text(_gap(0, caption=f"图 9-9 xx{SUFFIX}"), cfg).count(SUFFIX) == 1
    assert _caption_text(_gap(0, caption="图 9-9 xx"), cfg).endswith(SUFFIX)


def test_export_document_unknown_format_raises(tmp_path: Path):
    doc = Document(source_path="x", fmt="rtf")
    with pytest.raises(ValueError):
        export_document(doc, [], tmp_path / "o.rtf", Config())


def test_register_exporter_extension(tmp_path: Path):
    from docchart.insertion.exporters import _EXPORTERS, register_exporter
    calls = {}

    def _fake(doc, gaps, out_path, cfg):
        calls["ok"] = True
        Path(out_path).write_text("stub", encoding="utf-8")
        return Path(out_path)

    register_exporter("stubfmt", _fake)
    try:
        out = export_document(Document(source_path="x", fmt="stubfmt"),
                              [], tmp_path / "o.txt", Config())
        assert calls["ok"] and out.read_text(encoding="utf-8") == "stub"
    finally:
        _EXPORTERS.pop("stubfmt", None)
