# -*- coding: utf-8 -*-
"""端到端测试：检测 -> 生成 -> 插入 -> 导出 -> 校验 全流程。"""

from pathlib import Path

from docchart.config import Config
from docchart.pipeline import analyze, fill

ROOT = Path(__file__).resolve().parents[1]


def test_analyze_example_markdown():
    doc, gaps = analyze(ROOT / "examples" / "sample_report.md")
    kinds = {g.gap_type.value for g in gaps}
    # 示例文档刻意埋入全部五类信号
    assert {"dangling_reference", "caption_without_image",
            "empty_placeholder", "table_without_chart"} <= kinds
    assert all(g.section for g in gaps)


def test_fill_markdown_end_to_end(tmp_path: Path):
    src = tmp_path / "doc.md"
    src.write_text(
        "# 工程简报\n\n## 进度\n\n"
        "| 月 | 浇筑量(m³) |\n| --- | --- |\n| 3月 | 120 |\n| 4月 | 260 |\n| 5月 | 310 |\n\n"
        "后续拆除工艺流程如下图所示：\n\n检查平台→拆除连墙件→拆除架体→材料转运\n",
        encoding="utf-8")
    out = tmp_path / "doc_filled.md"
    report = fill(src, out_path=out, cfg=Config())
    assert report["summary"]["filled"] >= 2
    text = out.read_text(encoding="utf-8")
    # 表格图插在表格后、流程插图插在引用句后，且不破坏原内容
    assert "| 5月 | 310 |" in text and "autofill_" in text and "# 工程简报" in text
    tables_pos = text.find("| 5月 | 310 |")
    first_img = text.find("autofill_")
    assert first_img > tables_pos
    assert report["verification"]["passed"]
    for g in report["gaps"]:
        if g["status"] == "filled":
            assert Path(g["output_image"]).exists()


def test_fill_markdown_skip_reason(tmp_path: Path):
    src = tmp_path / "s.md"
    src.write_text("# A\n\n本段仅引用图 8-8 且全文无任何数据可提取。\n", encoding="utf-8")
    report = fill(src, out_path=tmp_path / "s_out.md", cfg=Config())
    skipped = [g for g in report["gaps"] if g["status"] == "skipped"]
    assert skipped and "建议补充" in skipped[0]["skip_reason"]


def test_fill_html_end_to_end(tmp_path: Path):
    src = tmp_path / "p.html"
    src.write_text(
        "<html><body><h1>塔吊方案</h1><p>拆除流程如下图所示：验收→固定→顶升</p>"
        "</body></html>", encoding="utf-8")
    out = tmp_path / "p_out.html"
    report = fill(src, out_path=out, cfg=Config())
    assert report["summary"]["filled"] >= 1
    text = out.read_text(encoding="utf-8")
    assert '<figure class="auto-filled">' in text and "<h1>塔吊方案</h1>" in text
    assert report["verification"]["passed"]


def test_fill_docx_end_to_end(tmp_path: Path):
    from docx import Document as Docx
    src = tmp_path / "t.docx"
    d = Docx()
    d.add_heading("第一章 概况", level=1)
    d.add_paragraph("材料使用量统计如下，控制成本变化趋势。")
    t = d.add_table(rows=4, cols=2)
    for i, (a, b) in enumerate([("月份", "钢管(t)"), ("3月", "12"), ("4月", "26"), ("5月", "18")]):
        t.cell(i, 0).text = a
        t.cell(i, 1).text = b
    d.add_paragraph("脚手架拆除工艺流程如下图所示：安全网→脚手板→连墙件→立杆")
    d.save(src)

    out = tmp_path / "t_filled.docx"
    report = fill(src, out_path=out, cfg=Config())
    assert report["summary"]["filled"] >= 2, report["gaps"]
    # 用 python-docx 重新打开导出文档验证：内嵌图片数、原内容保留
    from docx.shared import Inches  # noqa: F401
    d2 = Docx(str(out))
    imgs = d2.inline_shapes
    assert len(imgs) >= 2
    all_text = "\n".join(p.text for p in d2.paragraphs)
    assert "材料使用量统计" in all_text           # 原文未丢
    assert "（AI 补全）" in all_text              # 图注已写入
    assert report["verification"]["passed"]
    # 原文件未被修改
    d_orig = Docx(str(src))
    assert len(d_orig.inline_shapes) == 0
