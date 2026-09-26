# -*- coding: utf-8 -*-
"""解析器单元测试：结构识别、行号/偏移、图注与图片。"""

from pathlib import Path

from docchart.models import BlockKind
from docchart.parsers import parse_file, supported_formats


def test_supported_formats():
    assert {".md", ".html", ".docx"} <= set(supported_formats())


def test_markdown_parse(tmp_path: Path):
    md = tmp_path / "a.md"
    md.write_text(
        "# 标题一\n\n正文段落，含数值 1200 元。\n\n"
        "| 月 | 量 |\n| --- | --- |\n| 3月 | 10 |\n| 4月 | 20 |\n\n"
        "![某图](img/x.png)\n\n图 1-1 图注标题\n\n- 列表项\n",
        encoding="utf-8")
    doc = parse_file(md)
    kinds = [b.kind for b in doc.blocks]
    assert kinds[0] == BlockKind.HEADING and doc.blocks[0].text == "标题一"
    assert BlockKind.TABLE in kinds
    table = next(b for b in doc.blocks if b.kind == BlockKind.TABLE)
    assert table.rows[0] == ["月", "量"] and len(table.rows) == 3
    assert BlockKind.IMAGE in kinds
    assert BlockKind.CAPTION in kinds
    cap = next(b for b in doc.blocks if b.kind == BlockKind.CAPTION)
    assert cap.src["line_start"] == 11  # 图注所在行号（0-based）
    assert doc.blocks[0].section == "标题一"


def test_html_parse_offsets(tmp_path: Path):
    html = tmp_path / "a.html"
    src = "<h1>标题</h1><p>引用如下图所示。</p><table><tr><th>列</th></tr><tr><td>1</td></tr></table>"
    html.write_text(src, encoding="utf-8")
    doc = parse_file(html)
    kinds = [b.kind for b in doc.blocks]
    assert kinds == [BlockKind.HEADING, BlockKind.PARAGRAPH, BlockKind.TABLE]
    # 字符偏移必须能回指原文，导出回插依赖该性质
    p = doc.blocks[1]
    assert src[p.src["html_start"]:p.src["html_end"]] == "<p>引用如下图所示。</p>"
    assert doc.blocks[2].rows == [["列"], ["1"]]


def test_docx_parse(tmp_path: Path):
    from docx import Document as Docx
    f = tmp_path / "a.docx"
    d = Docx()
    d.add_heading("第一章 概况", level=1)
    d.add_paragraph("立杆步距如下图所示：")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "层"
    t.cell(0, 1).text = "荷载(kN)"
    t.cell(1, 0).text = "2"
    t.cell(1, 1).text = "3.2"
    d.add_paragraph("图 1-1 节点示意")
    d.save(f)
    doc = parse_file(f)
    kinds = [b.kind for b in doc.blocks]
    assert BlockKind.HEADING in kinds
    assert BlockKind.TABLE in kinds
    assert BlockKind.CAPTION in kinds
    table = next(b for b in doc.blocks if b.kind == BlockKind.TABLE)
    assert table.rows[1] == ["2", "3.2"]
    assert doc.native is not None  # 原生对象保留，供回插


def test_unsupported_ext(tmp_path: Path):
    f = tmp_path / "a.pdf"
    f.write_bytes(b"%PDF")
    try:
        parse_file(f)
        assert False, "应抛出不支持异常"
    except ValueError as e:
        assert "PDF" in str(e)
