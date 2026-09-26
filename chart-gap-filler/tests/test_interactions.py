# -*- coding: utf-8 -*-
"""跨组件交互与修复回归测试：解析->检测->生成->导出->校验 的数据契约。

同时锁定本轮修复的行为：
- 图注句式误判（"图4所示的…"不得再被当作图注）；
- R6 图号断裂只报告不出图；
- fill 恒定产出导出文档、导出容器格式守卫；
- DOCX 场景 svg 配置强制降级 png；
- table_min_numeric_cols / caption_prefix 配置生效；
- 包级公共 API 与解析器/导出器扩展点。
"""

import copy
import json
from pathlib import Path

import pytest

from docchart.config import Config, DEFAULT_CONFIG
from docchart.detection import Detector
from docchart.generation.recommender import build_spec
from docchart.models import Block, BlockKind, ChartData, ChartGap, GapType, Series, figure_caption_of
from docchart.parsers import parse_file, register_parser
from docchart.pipeline import fill

ROOT = Path(__file__).resolve().parents[1]


def _analyze_md(tmp_path: Path, text: str, cfg: Config | None = None):
    src = tmp_path / "t.md"
    src.write_text(text, encoding="utf-8")
    doc = parse_file(src)
    return doc, Detector(doc, cfg or Config()).detect()


# ---------------------------------------------------------------------------
# 图注识别回归（BUG：引用句式误判为图注）
# ---------------------------------------------------------------------------

def test_figure_caption_of_reference_phrases():
    assert figure_caption_of("图4所示的脚手架必须验收后使用。") is None
    assert figure_caption_of("图4层结构验算如下") is None
    assert figure_caption_of("图 1-1 节点示意") == ("1-1", "节点示意")
    assert figure_caption_of("图4-1荷载统计") == ("4-1", "荷载统计")
    assert figure_caption_of("图 4 进度计划") == ("4", "进度计划")


def test_markdown_reference_line_not_caption_block(tmp_path: Path):
    doc, gaps = _analyze_md(
        tmp_path, "# A\n\n图4所示的架体搭设应分层分段进行，每段高度不超过24米，经验收合格后方可继续。\n")
    assert [b.kind for b in doc.blocks] == [BlockKind.HEADING, BlockKind.PARAGRAPH]
    # 只按悬空引用报告，不再伪造"图注无图"缺口去错误插图
    assert {g.gap_type for g in gaps} == {GapType.DANGLING_REF}


def test_docx_reference_line_not_caption(tmp_path: Path):
    from docx import Document as Docx
    src = tmp_path / "a.docx"
    d = Docx()
    d.add_paragraph("图4所示的架体搭设应分层分段进行，验收后使用。")
    d.save(src)
    doc = parse_file(src)
    assert doc.blocks[0].kind == BlockKind.PARAGRAPH


# ---------------------------------------------------------------------------
# R6 图号断裂：检测 -> 报告 -> 不强行出图
# ---------------------------------------------------------------------------

def test_broken_numbering_detected_and_skipped(tmp_path: Path):
    src = tmp_path / "n.md"
    src.write_text("# A\n\n图 1-1 aaa\n\n![](x.png)\n\n图 1-3 ccc\n\n![](y.png)\n",
                   encoding="utf-8")
    out = tmp_path / "n_out.md"
    report = fill(src, out_path=out, cfg=Config())
    broken = [g for g in report["gaps"] if g["gap_type"] == "broken_figure_number"]
    assert broken, report["gaps"]
    assert broken[0]["status"] == "skipped"
    assert "人工核对" in broken[0]["skip_reason"]
    # 导出文档保持原样，未插入臆测的图
    assert "![autofill" not in out.read_text(encoding="utf-8")


def test_numbering_continuous_no_report(tmp_path: Path):
    _, gaps = _analyze_md(
        tmp_path, "# A\n\n图 1-1 aaa\n\n![](x.png)\n\n图 1-2 bbb\n\n![](y.png)\n")
    assert not [g for g in gaps if g.gap_type == GapType.BROKEN_NUMBERING]


# ---------------------------------------------------------------------------
# 导出链路与报告契约
# ---------------------------------------------------------------------------

def test_fill_without_gaps_still_exports(tmp_path: Path):
    src = tmp_path / "clean.md"
    text = "# A\n\n纯文本段落，无任何配图信号。\n"
    src.write_text(text, encoding="utf-8")
    report = fill(src, out_path=tmp_path / "clean_out.md", cfg=Config())
    out = Path(report["exported_document"])
    assert out.exists() and out.read_text(encoding="utf-8") == text
    assert report["summary"]["total_gaps"] == 0
    assert report["verification"]["passed"]


def test_fill_rejects_container_mismatch(tmp_path: Path):
    src = tmp_path / "a.md"
    src.write_text("# A\n\n正文。\n", encoding="utf-8")
    with pytest.raises(ValueError):
        fill(src, out_path=tmp_path / "a.docx", cfg=Config())


def test_report_schema_contract(tmp_path: Path):
    src = tmp_path / "s.md"
    src.write_text(
        "# 文档\n\n| 月 | 量 |\n| --- | --- |\n| 3月 | 10 |\n| 4月 | 20 |\n| 5月 | 30 |\n",
        encoding="utf-8")
    report = fill(src, out_path=tmp_path / "s_out.md", cfg=Config())
    assert set(report) == {"source", "format", "blocks", "exported_document",
                           "assets_dir", "summary", "gaps", "verification"}
    assert set(report["summary"]) == {"total_gaps", "filled", "skipped"}
    g0 = report["gaps"][0]
    assert {"anchor_block_id", "gap_type", "confidence", "status",
            "caption", "section", "skip_reason", "output_image"} <= set(g0)
    # 报告必须可 JSON 序列化（下游 CLI/服务以此为数据交换格式）
    json.dumps(report, ensure_ascii=False)


def test_fill_docx_forces_png_when_svg_configured(tmp_path: Path):
    from docx import Document as Docx
    src = tmp_path / "t.docx"
    d = Docx()
    d.add_heading("第一章 概况", 1)
    d.add_paragraph("材料用量统计，控制成本变化趋势。")
    t = d.add_table(rows=4, cols=2)
    for i, (a, b) in enumerate([("月份", "钢管(t)"), ("3月", "12"), ("4月", "26"), ("5月", "18")]):
        t.cell(i, 0).text, t.cell(i, 1).text = a, b
    d.save(src)
    cfg = Config()
    cfg.data["output"]["image_format"] = "svg"
    report = fill(src, out_path=tmp_path / "t_out.docx", cfg=cfg)
    filled = [g for g in report["gaps"] if g["status"] == "filled"]
    assert filled and all(Path(g["output_image"]).suffix == ".png" for g in filled)
    assert report["verification"]["passed"]


# ---------------------------------------------------------------------------
# 配置接线
# ---------------------------------------------------------------------------

def test_table_min_numeric_cols_wired(tmp_path: Path):
    md = ("# A\n\n| 项目 | 计划 | 实际 | 备注 |\n| --- | --- | --- | --- |\n"
          "| a | 10 | 12 | 正常 |\n| b | 20 | 22 | 正常 |\n| c | 30 | 31 | 正常 |\n")
    _, gaps = _analyze_md(tmp_path, md)
    assert [g for g in gaps if g.gap_type == GapType.TABLE_NO_CHART]
    cfg = Config()
    cfg.data["detection"]["table_min_numeric_cols"] = 3  # 只有 2 个数值列 -> 不报
    _, gaps2 = _analyze_md(tmp_path, md, cfg)
    assert not [g for g in gaps2 if g.gap_type == GapType.TABLE_NO_CHART]


def test_caption_prefix_wired_in_dangling_ref(tmp_path: Path):
    cfg = Config()
    cfg.data["output"]["caption_prefix"] = "FIG"
    _, gaps = _analyze_md(tmp_path, "# A\n\n施工进度如图 9-9 所示。\n", cfg)
    g = next(g for g in gaps if g.gap_type == GapType.DANGLING_REF)
    assert g.caption.startswith("FIG9-9")


def test_config_load_json_and_missing(tmp_path: Path):
    cfg_file = tmp_path / "c.json"
    cfg_file.write_text(json.dumps({"output": {"image_format": "svg"},
                                    "detection": {"min_confidence": 0.9}}),
                        encoding="utf-8")
    cfg = Config.load(cfg_file)
    assert cfg.get("output", "image_format") == "svg"
    assert cfg.get("detection", "min_confidence") == 0.9
    # 未覆盖的键保持默认（深层合并不吃掉兄弟键）
    assert cfg.get("render", "dpi") == DEFAULT_CONFIG["render"]["dpi"]
    assert Config.load(tmp_path / "nope.json").get("output", "image_format") == "png"
    assert Config.load(None).get("detection", "ref_window_blocks") == 12


def test_config_load_yaml_optional(tmp_path: Path):
    pytest.importorskip("yaml")
    f = tmp_path / "c.yaml"
    f.write_text("render:\n  dpi: 96\n", encoding="utf-8")
    assert Config.load(f).get("render", "dpi") == 96


# ---------------------------------------------------------------------------
# 生成器回退逻辑（BUG：以系列名充当流程步骤）
# ---------------------------------------------------------------------------

def _gap(suggested: str) -> ChartGap:
    return ChartGap(anchor_block_id=0, gap_type=GapType.SEMANTIC_HINT, reason="",
                    confidence=1.0, suggested_type=suggested, caption="标题")


def test_flow_spec_falls_back_to_categories_not_series_names():
    data = ChartData(categories=["开挖", "支护", "回填"],
                     series=[Series(name="进度", values=[1, 2, 3])], complete=False)
    spec = build_spec(_gap("flow"), data, "文档")
    assert spec is not None and spec.chart_type == "flow"
    assert spec.steps == ["开挖", "支护", "回填"]   # 不再是表头系列名 ["进度"]
    assert "供核对" in spec.note


def test_flow_spec_rejects_single_node():
    data = ChartData(categories=["唯一项"], series=[Series("v", [1, 2, 3])], complete=False)
    assert build_spec(_gap("flow"), data, "文档") is None


# ---------------------------------------------------------------------------
# 扩展点与数据不变量
# ---------------------------------------------------------------------------

def test_block_id_equals_index_all_formats(tmp_path: Path):
    md = tmp_path / "i.md"
    md.write_text((ROOT / "examples" / "sample_report.md").read_text(encoding="utf-8"),
                  encoding="utf-8")
    html = tmp_path / "i.html"
    html.write_text("<html><body><h1>t</h1><p>p</p><img src='a.png'></body></html>",
                    encoding="utf-8")
    docx = tmp_path / "i.docx"
    from docx import Document as Docx
    d = Docx()
    d.add_heading("h", 1)
    d.add_paragraph("p")
    d.save(docx)
    for f in (md, html, docx):
        doc = parse_file(f)
        assert all(b.id == i for i, b in enumerate(doc.blocks)), f


def test_register_parser_extension(tmp_path: Path):
    from docchart import parsers
    from docchart.models import Document
    register_parser(".fake", lambda p: Document(source_path=str(p), fmt="fake"))
    try:
        f = tmp_path / "x.fake"
        f.write_text("anything", encoding="utf-8")
        assert parse_file(f).fmt == "fake"
    finally:
        parsers._REGISTRY.pop(".fake", None)


def test_public_api_surface():
    import docchart
    assert callable(docchart.analyze) and callable(docchart.fill)
    assert callable(docchart.save_report) and callable(docchart.supported_formats)
    with pytest.raises(AttributeError):
        docchart.no_such_symbol


def test_models_import_light():
    """models 是组件间数据契约，不应依赖渲染/解析重库。"""
    import subprocess
    import sys
    code = ("import sys; import docchart.models as m;"
            "print(any(x in sys.modules for x in ('matplotlib','docx')))")
    out = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True,
                         cwd=str(ROOT), timeout=60)
    assert out.stdout.strip().endswith("False"), out.stdout + out.stderr


def test_datasource_block_contract():
    """datasource 直接消费解析器产出的 Block.rows，行对齐是硬契约。"""
    from docchart.generation.datasource import extract_from_table
    b = Block(id=3, kind=BlockKind.TABLE,
              rows=[["月", "量", "备"], ["3月", "10", ""], ["4月", "20", "ok"]])
    data = extract_from_table(b)
    assert data and data.source_block_id == 3
    assert data.categories == ["3月", "4月"]
    assert [s.name for s in data.series] == ["量"]      # 备注列非数值被剔除
    assert copy.deepcopy(data).series[0].values == [10.0, 20.0]
