# -*- coding: utf-8 -*-
"""F2 段落级端到端验证：dangling 残式 -> 占位框 + 无双空格。"""
import sys
sys.path.insert(0, r"j:\编程\专项方案工具箱\backend")

import docx
from app.routers import export

cases = [
    "立杆轴向压力设计值 $N = $由支架立杆轴向力计算确定。",
    "式中 $N$ 为脚手架立杆轴向压力设计值。",
    "甲醛释放量控制应满足不超过 $L = $ 的目标值",
]

ok = True
for c in cases:
    p = docx.Document().add_paragraph()
    export._add_runs_with_inline_format(p, c)
    xml = p._p.xml
    texts = [r.text for r in p.runs]
    joined = "".join(texts)
    has_box = "<m:box>" in xml
    has_tian = ">待填<" in xml
    double_space = any("  " in t for t in texts)
    print("原文:", c)
    print("  run文本:", texts)
    print("  含占位框:", has_box, " 含待填:", has_tian, " run内双空格:", double_space)
    if "=" in c and "$" in c:
        if not (has_box and has_tian):
            ok = False
            print("  [FAIL] dangling 残式应生成占位框")
    print()

print("结果:", "PASS" if ok else "FAIL")
sys.exit(0 if ok else 1)
